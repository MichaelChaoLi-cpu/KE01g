#!/usr/bin/env python3
"""Shelter-Level Provisional Toilet Rebalancing Plan.

Plan: Report the complete 38-shelter Yatsushiro donor, recipient, transfer,
post-transfer inventory, residual-pressure, and evidence-qualification record
under full reported-surplus mobility.
Framework: Section 5 conditional mitigation contrast; Section 6 donor capacity
D, integer transfer x, post-transfer inventory, residual shortfall U, donor
protection, and priority-constrained nearest-donor rule; Section 7 conditional
rebalancing workflow. Movements are a best-case sensitivity, not dispatch.
"""

from __future__ import annotations

from pathlib import Path
import re

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

import figure_provisional_emergency_toilet_rebalancing_plan as rebalancing
from table_shelter_level_equity_demand_scenario_estimates import (
    AREA_LABELS,
    ENGLISH_SHELTER_NAMES,
    LOCATION_LABELS,
)


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/processed/shelter_equity_scenarios_preprocessed.parquet"
OUTPUT = (
    ROOT
    / "data/results/tables/"
    "Table_shelter_level_provisional_toilet_rebalancing_plan.xlsx"
)

JAPANESE_PATTERN = re.compile(r"[\u3040-\u30ff\u3400-\u9fff]")
WATER_LABELS = {
    "〇": "Available",
    "○": "Available",
    "△": "Partially available",
    "×": "Unavailable",
}
HEADERS = (
    "Stable Shelter ID",
    "Shelter Name",
    "Area",
    "Evacuees",
    "Water Status",
    "Temporary Toilets Installed",
    "Toilet Cars",
    "Prolonged Requirement",
    "Current Balance (+ Need / - Surplus)",
    "Plan Role",
    "Proposed Transfer In",
    "Proposed Transfer Out",
    "Transfer Partners (Units)",
    "Transfer Distance (km)",
    "Post-Transfer Inventory",
    "Residual Shortfall",
    "Donor Protection Rule",
    "Location Resolution",
)
MAIN_HEADERS = (
    "Stable Shelter ID",
    "Shelter Name",
    "Area",
    "Evacuees",
    "Water Status",
    "Reported Units (Temporary / Cars)",
    "Prolonged Requirement",
    "Current Balance (+ Need / - Surplus)",
    "Proposed Action (+ Receive / - Send)",
    "Transfer Links (Partner; Units; Distance)",
    "Post-Transfer Inventory",
    "Location Resolution",
)


def build_plan(frame: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    required = {
        "Municipality",
        "Shelter Number",
        "Shelter Name",
        "District",
        "Latitude",
        "Longitude",
        "Location Resolution",
        "Evacuees",
        "Water Status",
        "Temporary Toilets Installed",
        "Toilet Cars",
        "Scenario",
        "Estimated Functional Support Evacuees",
        "Estimated Female Functional Support Evacuees",
    }
    missing = required.difference(frame.columns)
    if missing:
        raise KeyError(f"Missing required columns: {sorted(missing)}")
    base = frame.loc[
        frame["Municipality"].eq("八代市")
        & frame["Scenario"].astype("string").eq("base")
    ].copy().reset_index(drop=True)
    if len(base) != 38 or base[list(required.difference({"Municipality", "Scenario"}))].isna().any().any():
        raise ValueError("Expected 38 complete Yatsushiro base-scenario records")
    screen = rebalancing.construct_screen(base)
    flows, post = rebalancing.full_mobility_flows(screen)
    if int(screen["Screening Shortfall"].sum()) != 31:
        raise ValueError("Expected 31 current screening-shortfall units")
    if int(screen["Reported Surplus"].sum()) != 72:
        raise ValueError("Expected 72 reported surplus units")
    if int(flows["Units"].sum()) != 31 or int(post["Residual Shortfall"].sum()) != 0:
        raise ValueError("Full mobility must move 31 units and resolve the screen")
    if int(post["Transfer In"].sum()) != int(post["Transfer Out"].sum()):
        raise ValueError("Transfer conservation failed")
    donors = post.loc[post["Transfer Out"].gt(0)]
    if (donors["Post-Transfer Inventory"] < donors["Prolonged Requirement"]).any():
        raise ValueError("A donor falls below its protected prolonged requirement")

    post["Stable Shelter ID"] = "Y" + post["Shelter Number"].astype(int).astype(str).str.zfill(2)
    index_to_id = post["Stable Shelter ID"].to_dict()
    partner_labels: dict[int, list[str]] = {int(index): [] for index in post.index}
    distance_values: dict[int, list[float]] = {int(index): [] for index in post.index}
    flow_records = []
    for _, flow in flows.iterrows():
        donor_index = int(flow["Donor Index"])
        recipient_index = int(flow["Recipient Index"])
        units = int(flow["Units"])
        distance = float(flow["Distance km"])
        donor_id = index_to_id[donor_index]
        recipient_id = index_to_id[recipient_index]
        partner_labels[donor_index].append(f"{recipient_id} ({units})")
        partner_labels[recipient_index].append(f"{donor_id} ({units})")
        distance_values[donor_index].append(distance)
        distance_values[recipient_index].append(distance)
        flow_records.append(
            {
                "Donor Shelter ID": donor_id,
                "Recipient Shelter ID": recipient_id,
                "Units": units,
                "Distance (km)": distance,
                "Donor Location Resolution": LOCATION_LABELS[str(post.loc[donor_index, "Location Resolution"])],
                "Recipient Location Resolution": LOCATION_LABELS[str(post.loc[recipient_index, "Location Resolution"])],
            }
        )
    flow_table = pd.DataFrame(flow_records)
    if len(flow_table) != 18:
        raise ValueError("Expected 18 donor-recipient movements")

    plan = pd.DataFrame()
    plan["Stable Shelter ID"] = post["Stable Shelter ID"]
    plan["Shelter Name"] = post["Stable Shelter ID"].map(ENGLISH_SHELTER_NAMES)
    plan["Area"] = post["District"].map(AREA_LABELS)
    plan["Evacuees"] = post["Evacuees"].astype(int)
    plan["Water Status"] = post["Water Status"].astype(str).map(WATER_LABELS)
    plan["Temporary Toilets Installed"] = post["Temporary Toilets Installed"].astype(int)
    plan["Toilet Cars"] = post["Toilet Cars"].astype(int)
    plan["Prolonged Requirement"] = post["Prolonged Requirement"].astype(int)
    plan["Current Balance (+ Need / - Surplus)"] = (
        post["Screening Shortfall"] - post["Reported Surplus"]
    ).astype(int)
    plan["Plan Role"] = "No transfer"
    plan.loc[post["Transfer Out"].gt(0), "Plan Role"] = "Donor"
    plan.loc[post["Transfer In"].gt(0), "Plan Role"] = "Recipient"
    plan["Proposed Transfer In"] = post["Transfer In"].astype(int)
    plan["Proposed Transfer Out"] = post["Transfer Out"].astype(int)
    plan["Transfer Partners (Units)"] = [
        "; ".join(partner_labels[int(index)]) if partner_labels[int(index)] else "N/A"
        for index in post.index
    ]
    plan["Transfer Distance (km)"] = [
        (
            "N/A"
            if not distance_values[int(index)]
            else f"{min(distance_values[int(index)]):.2f}"
            if max(distance_values[int(index)]) - min(distance_values[int(index)]) < 0.005
            else f"{min(distance_values[int(index)]):.2f}-{max(distance_values[int(index)]):.2f}"
        )
        for index in post.index
    ]
    plan["Post-Transfer Inventory"] = post["Post-Transfer Inventory"].astype(int)
    plan["Residual Shortfall"] = post["Residual Shortfall"].astype(int)
    plan["Donor Protection Rule"] = [
        f"Keep >= {int(requirement)}"
        if transfer_out > 0
        else "N/A"
        for requirement, transfer_out in zip(
            post["Prolonged Requirement"], post["Transfer Out"], strict=True
        )
    ]
    plan["Location Resolution"] = post["Location Resolution"].astype(str).map(LOCATION_LABELS)
    if plan[list(HEADERS)].isna().any().any():
        raise ValueError("English rebalancing plan contains missing values")
    role_order = plan["Plan Role"].map({"Recipient": 0, "Donor": 1, "No transfer": 2})
    plan = (
        plan.assign(_role_order=role_order, _tier=post["Verification Tier"].to_numpy())
        .sort_values(
            ["_role_order", "_tier", "Current Balance (+ Need / - Surplus)", "Stable Shelter ID"],
            ascending=[True, True, False, True],
            kind="stable",
        )
        .drop(columns=["_role_order", "_tier"])
        .reset_index(drop=True)
    )
    return plan[list(HEADERS)], flow_table


def write_workbook(plan: pd.DataFrame, flows: pd.DataFrame) -> None:
    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    sheet = workbook.active
    sheet.title = "Rebalancing Plan"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "D2"
    sheet.sheet_view.zoomScale = 80
    sheet.append(list(MAIN_HEADERS))
    for row_number, (_, row) in enumerate(plan.iterrows(), start=2):
        sheet.append(
            [
                row["Stable Shelter ID"],
                f"='Full Rebalancing Record'!B{row_number}",
                f"='Full Rebalancing Record'!C{row_number}",
                f"='Full Rebalancing Record'!D{row_number}",
                f"='Full Rebalancing Record'!E{row_number}",
                (
                    f'=TEXT(\'Full Rebalancing Record\'!F{row_number},"0")&" / "&'
                    f'TEXT(\'Full Rebalancing Record\'!G{row_number},"0")'
                ),
                f"='Full Rebalancing Record'!H{row_number}",
                f"='Full Rebalancing Record'!I{row_number}",
                (
                    f"='Full Rebalancing Record'!K{row_number}-"
                    f"'Full Rebalancing Record'!L{row_number}"
                ),
                (
                    f'=IF(\'Full Rebalancing Record\'!M{row_number}="N/A","N/A",'
                    f'\'Full Rebalancing Record\'!M{row_number}&"; "&'
                    f'\'Full Rebalancing Record\'!N{row_number}&" km")'
                ),
                f"='Full Rebalancing Record'!O{row_number}",
                f"='Full Rebalancing Record'!R{row_number}",
            ]
        )

    last_row = sheet.max_row
    excel_table = Table(displayName="CompactRebalancingPlan", ref=f"A1:L{last_row}")
    excel_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(excel_table)
    navy = "27445C"
    text = "263746"
    border = Side(style="thin", color="D8DEE3")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=8, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="9EB1C0"))
    sheet.row_dimensions[1].height = 72
    for row in sheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=12):
        for column_index, cell in enumerate(row):
            cell.font = Font(name="Aptos", size=7.5, color=text)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(
                left=border if column_index > 0 else Side(style=None), bottom=border
            )
        for index in (3, 6, 7, 8, 10):
            row[index].alignment = Alignment(horizontal="right", vertical="center")
        sheet.row_dimensions[row[0].row].height = 34
    for column in ("D", "G", "H", "I", "K"):
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0;-#,##0;0"

    red_fill = PatternFill("solid", fgColor="F3CEC5")
    green_fill = PatternFill("solid", fgColor="CDE2D7")
    gray_fill = PatternFill("solid", fgColor="E7EBEE")
    sheet.conditional_formatting.add(
        f"H2:H{last_row}", CellIsRule(operator="greaterThan", formula=["0"], fill=red_fill)
    )
    sheet.conditional_formatting.add(
        f"H2:H{last_row}", CellIsRule(operator="lessThan", formula=["0"], fill=green_fill)
    )
    sheet.conditional_formatting.add(
        f"H2:H{last_row}", CellIsRule(operator="equal", formula=["0"], fill=gray_fill)
    )
    sheet.conditional_formatting.add(
        f"I2:I{last_row}", CellIsRule(operator="greaterThan", formula=["0"], fill=red_fill)
    )
    sheet.conditional_formatting.add(
        f"I2:I{last_row}", CellIsRule(operator="lessThan", formula=["0"], fill=green_fill)
    )
    sheet.conditional_formatting.add(
        f"I2:I{last_row}", CellIsRule(operator="equal", formula=["0"], fill=gray_fill)
    )
    widths = {
        "A": 13,
        "B": 37,
        "C": 14,
        "D": 10,
        "E": 19,
        "F": 21,
        "G": 17,
        "H": 20,
        "I": 20,
        "J": 34,
        "K": 18,
        "L": 23,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet["H1"].comment = Comment(
        "Positive values are required increases; negative values are reported units above the shelter's prolonged benchmark requirement. Neither is a functional-capacity measure.",
        "OpenAI",
    )
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 2
    sheet.print_title_rows = "1:1"
    sheet.print_area = f"A1:L{last_row}"

    inputs = workbook.create_sheet("Plan Inputs")
    inputs.title = "Full Rebalancing Record"
    inputs.sheet_view.showGridLines = False
    inputs.append(list(HEADERS))
    for row_number, (_, row) in enumerate(plan.iterrows(), start=2):
        values = row.tolist()
        values[14] = f"=F{row_number}+K{row_number}-L{row_number}"
        values[15] = f"=MAX(0,H{row_number}-O{row_number})"
        inputs.append(values)
    for cell in inputs[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=8, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    inputs.row_dimensions[1].height = 68
    for row in inputs.iter_rows(min_row=2, max_row=inputs.max_row):
        for cell in row:
            cell.font = Font(name="Aptos", size=7.5, color=text)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=border)
        inputs.row_dimensions[row[0].row].height = 32
    for column, width in widths.items():
        if column in {"A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L"}:
            inputs.column_dimensions[column].width = width
    for column, width in {"M": 23, "N": 16, "O": 16, "P": 13, "Q": 17, "R": 22}.items():
        inputs.column_dimensions[column].width = width
    inputs.freeze_panes = "D2"
    inputs.auto_filter.ref = f"A1:R{inputs.max_row}"
    full_table = Table(displayName="FullRebalancingRecord", ref=f"A1:R{inputs.max_row}")
    full_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    inputs.add_table(full_table)

    flow_sheet = workbook.create_sheet("Transfer Flows")
    flow_sheet.sheet_view.showGridLines = False
    flow_headers = list(flows.columns)
    flow_sheet.append(flow_headers)
    for _, row in flows.iterrows():
        flow_sheet.append(row.tolist())
    for cell in flow_sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=8.5, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    flow_sheet.row_dimensions[1].height = 42
    for row in flow_sheet.iter_rows(min_row=2, max_row=flow_sheet.max_row):
        for cell in row:
            cell.font = Font(name="Aptos", size=8.5, color=text)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=border)
        row[2].alignment = Alignment(horizontal="right", vertical="center")
        row[3].alignment = Alignment(horizontal="right", vertical="center")
        row[3].number_format = "0.00"
    for column, width in {"A": 19, "B": 21, "C": 10, "D": 14, "E": 25, "F": 28}.items():
        flow_sheet.column_dimensions[column].width = width
    flow_sheet.freeze_panes = "A2"
    flow_sheet.auto_filter.ref = f"A1:F{flow_sheet.max_row}"

    notes = workbook.create_sheet("Notes")
    notes.sheet_view.showGridLines = False
    notes.column_dimensions["A"].width = 31
    notes.column_dimensions["B"].width = 108
    notes.append(["Field", "Definition"])
    note_rows = [
        ("Full-mobility scenario", "Every reported temporary toilet above a donor shelter's prolonged benchmark requirement is assumed serviceable, movable, accessible by road, and compatible with setup and waste-service conditions."),
        ("Recipient order", "Ascending verification tier, then descending current shortfall, Base functional-support demand, Base female functional-support demand, and evacuees; stable shelter ID breaks ties."),
        ("Donor selection", "For each recipient in priority order, use the eligible donor with the shortest great-circle distance; stable shelter ID breaks distance ties."),
        ("Donor protection", "No donor may fall below its own sitewise prolonged benchmark requirement after transfer."),
        ("Residual result", "Residual screening shortfall is zero at all 38 shelters under this full-mobility sensitivity; the complete row-level field remains in Full Rebalancing Record."),
        ("Current balance", "Positive = prolonged benchmark requirement above reported temporary toilets; negative = reported temporary toilets above the benchmark. Fixed stalls and toilet cars are excluded."),
        ("Location evidence", "Exact shelter matches support point-level distance screening. District-anchor fallbacks do not represent verified facility coordinates or dispatch-ready routes."),
        ("Interpretation boundary", "The plan is a best-case spatial-rebalancing sensitivity. It is not a verified dispatch order, functional-capacity finding, accessible-compliance result, or global transport optimum."),
    ]
    for row in note_rows:
        notes.append(row)
    for cell in notes[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
    for row in notes.iter_rows(min_row=2, max_row=notes.max_row):
        row[0].font = Font(name="Aptos", size=9, bold=True, color=text)
        row[1].font = Font(name="Aptos", size=9, color=text)
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = Border(bottom=border)
        notes.row_dimensions[row[0].row].height = 48
    notes.freeze_panes = "A2"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)


def validate_output(plan: pd.DataFrame, flows: pd.DataFrame) -> None:
    workbook = load_workbook(OUTPUT, data_only=False)
    if workbook.sheetnames != ["Rebalancing Plan", "Full Rebalancing Record", "Transfer Flows", "Notes"]:
        raise ValueError("Unexpected workbook sheet structure")
    sheet = workbook["Rebalancing Plan"]
    if sheet.max_row != 39 or sheet.max_column != 12:
        raise ValueError("Unexpected main-table dimensions")
    formulas = [
        cell
        for workbook_sheet in workbook.worksheets
        for row in workbook_sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    if len(formulas) != 494:
        raise ValueError(f"Unexpected formula count: {len(formulas)}")
    japanese_cells = []
    for workbook_sheet in workbook.worksheets:
        for row in workbook_sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and JAPANESE_PATTERN.search(cell.value):
                    japanese_cells.append(f"{workbook_sheet.title}!{cell.coordinate}")
    if japanese_cells:
        raise ValueError(f"Japanese text remains in workbook: {japanese_cells[:5]}")
    if int(plan["Proposed Transfer In"].sum()) != 31:
        raise ValueError("Plan inflow total changed")
    if int(plan["Proposed Transfer Out"].sum()) != 31:
        raise ValueError("Plan outflow total changed")
    if int(plan["Residual Shortfall"].sum()) != 0 or int(flows["Units"].sum()) != 31:
        raise ValueError("Full-mobility accounting changed")


def main() -> None:
    plan, flows = build_plan(pd.read_parquet(SOURCE))
    write_workbook(plan, flows)
    validate_output(plan, flows)
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print("Main table: 38 shelters x 12 columns; full 18-column record retained")
    print(
        f"Recipients: {int(plan['Plan Role'].eq('Recipient').sum())}; "
        f"donors: {int(plan['Plan Role'].eq('Donor').sum())}; "
        f"flow records: {len(flows)}"
    )
    print("Transferred units: 31 in = 31 out; residual screening shortfall: 0")


if __name__ == "__main__":
    main()
