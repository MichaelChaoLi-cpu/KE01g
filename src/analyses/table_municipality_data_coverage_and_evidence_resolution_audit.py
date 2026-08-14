#!/usr/bin/env python3
"""Municipality Data Coverage and Evidence Resolution Audit.

Plan: Audit the 11 municipalities reporting open shelters for the availability
and temporal resolution of municipality totals, shelter rosters and occupancy,
water and deployment records, residential sex inputs, and sex-specific
functional-support proxy inputs.
Framework: Section 5's two-tier evidence design; Section 6's municipality
scenario inputs; and Section 7's prerequisite audit of common-snapshot totals,
sex reconciliation, reference-period alignment, and permitted analysis level.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]
COVERAGE_SOURCE = (
    ROOT / "data/exp/data-source-expansion/municipality_evidence_coverage.csv"
)
MUNICIPALITY_SOURCE = (
    ROOT
    / "data/processed/"
    "kumamoto_prefecture_municipality_shelter_totals_2026-08-12_1400_preprocessed.parquet"
)
LTC_TOTAL_SOURCE = ROOT / "data/processed/mhlw_ltc_2026_04_total_preprocessed.parquet"
LTC_MALE_SOURCE = ROOT / "data/processed/mhlw_ltc_2026_04_male_preprocessed.parquet"
LTC_FEMALE_SOURCE = ROOT / "data/processed/mhlw_ltc_2026_04_female_preprocessed.parquet"
OUTPUT = (
    ROOT
    / "data/results/tables/"
    "Table_municipality_data_coverage_and_evidence_resolution_audit.xlsx"
)

HEADERS = (
    "Municipality",
    "Municipality Observation Time",
    "Open Shelter Count",
    "Evacuees",
    "Shelter Roster Availability",
    "Shelter Occupancy Availability",
    "Water Status Availability",
    "Deployment Records Availability",
    "Residential Sex Input Status",
    "Functional Support Input Status",
    "Certification Sex Reconciliation Status",
    "Municipality Evidence Tier",
    "Permitted Analysis Level",
    "Reference-Period Alignment",
    "Decisive Limitation",
)

LEVEL_COLUMNS = (
    "Support Level 1",
    "Support Level 2",
    "Care Level 1",
    "Care Level 2",
    "Care Level 3",
    "Care Level 4",
    "Care Level 5",
    "Certified Total",
)

STATUS_MAPS = {
    "shelter_roster_status": {
        "available_near_time_2026-08-12_1200": "Available: 12:00 (-2 h)",
        "single_open_shelter_identity_supported": "Identity available: 1 shelter",
        "four_open_shelter_identities_supported": "Identities available: 4 shelters",
        "single_shelter_identity_partially_supported": "Partial identity: 1 shelter",
        "not_found_at_snapshot": "Not found for event snapshot",
    },
    "shelter_occupancy_status": {
        "available_near_time_2026-08-12_1200": "Available: 12:00 (-2 h)",
        "not_found_at_snapshot": "Not found for event snapshot",
    },
    "water_status_availability": {
        "available_near_time": "Available: 12:00 (-2 h)",
        "not_shelter_specific": "Available only as city context",
        "not_found": "Not found",
    },
    "deployment_records_availability": {
        "available": "Available: shelter and placement records",
        "partial": "Partial: selected current placements",
        "not_found": "Not found",
    },
    "evidence_tier": {
        "tier_2_shelter_detailed": "Tier 2: shelter-detailed",
        "tier_1_municipality_plus_identity": "Tier 1: municipality + identity",
        "tier_1_municipality_plus_partial_identity": (
            "Tier 1: municipality + partial identity"
        ),
        "tier_1_municipality_only": "Tier 1: municipality-only",
    },
    "permitted_analysis_level": {
        "municipality screening and shelter-level demand scenarios": (
            "Municipality screen + shelter demand scenarios"
        ),
        "municipality screening and shelter-level demand and deployment scenarios": (
            "Municipality screen + shelter demand/deployment scenarios"
        ),
        "municipality screening; shelter identity as context only": (
            "Municipality screen; shelter identity context only"
        ),
        "municipality screening; shelter identities as context only": (
            "Municipality screen; shelter identity context only"
        ),
        "municipality screening only": "Municipality screening only",
    },
}

LIMITATIONS = {
    "Kumamoto City": (
        "Water is not shelter-specific and deployment coverage is partial; "
        "fixed/functional stalls and subgroup occupants are unobserved."
    ),
    "Yatsushiro City": (
        "Shelter evidence is two hours earlier than the municipality total; "
        "fixed/functional stalls and subgroup occupants are unobserved."
    ),
    "Uto City": (
        "Only the open-shelter identity is supported; snapshot occupancy, water, "
        "and deployment evidence were not found."
    ),
    "Mifune Town": (
        "Only a partial shelter identity is supported; snapshot occupancy, water, "
        "and deployment evidence were not found."
    ),
    "Kosa Town": (
        "Open-shelter identities are supported, but snapshot occupancy, water, "
        "and deployment evidence were not found."
    ),
}

DEFAULT_LIMITATION = (
    "Event-snapshot shelter roster, occupancy, water, and deployment evidence "
    "were not found; only municipality screening is supported."
)


def read_and_validate_inputs() -> tuple[pd.DataFrame, pd.DataFrame]:
    coverage = pd.read_csv(COVERAGE_SOURCE, dtype={"municipality_code": "string"})
    municipality = pd.read_parquet(MUNICIPALITY_SOURCE)
    if len(coverage) != 11 or coverage["municipality_english"].nunique() != 11:
        raise ValueError("Expected one coverage record for each of 11 municipalities")
    if len(municipality) != 11 or municipality["Municipality"].nunique() != 11:
        raise ValueError("Expected one observed record for each of 11 municipalities")
    if municipality["Municipality Observation Time"].nunique() != 1:
        raise ValueError("Municipality records do not share one observation time")
    if int(municipality["Open Shelter Count"].sum()) != 81:
        raise ValueError("Municipality records do not reproduce 81 open shelters")
    if int(municipality["Evacuees"].sum()) != 3585:
        raise ValueError("Municipality records do not reproduce 3,585 evacuees")

    merged = coverage.merge(
        municipality[
            [
                "Municipality Code",
                "Municipality",
                "Municipality Observation Time",
                "Open Shelter Count",
                "Evacuees",
                "Municipality Evidence Tier",
            ]
        ],
        left_on="municipality_code",
        right_on="Municipality Code",
        how="left",
        validate="one_to_one",
        suffixes=("_audit", "_processed"),
    )
    if merged["Municipality"].isna().any():
        raise ValueError("Coverage audit did not match all processed municipalities")
    if not merged["municipality_english"].eq(merged["Municipality"]).all():
        raise ValueError("Municipality English-name mismatch")
    if not merged["open_shelters"].eq(merged["Open Shelter Count"]).all():
        raise ValueError("Open-shelter count mismatch between audit and processed data")
    if not merged["evacuees"].eq(merged["Evacuees"]).all():
        raise ValueError("Evacuee count mismatch between audit and processed data")

    sex_frames = []
    for sex, path in (
        ("Total", LTC_TOTAL_SOURCE),
        ("Male", LTC_MALE_SOURCE),
        ("Female", LTC_FEMALE_SOURCE),
    ):
        frame = pd.read_parquet(path, columns=["Municipality", *LEVEL_COLUMNS])
        frame = frame.loc[frame["Municipality"].isin(merged["municipality_japanese"])]
        frame = frame.rename(
            columns={
                "Municipality": "municipality_japanese",
                **{column: f"{sex} {column}" for column in LEVEL_COLUMNS},
            }
        )
        sex_frames.append(frame)

    audit_inputs = sex_frames[0]
    for frame in sex_frames[1:]:
        audit_inputs = audit_inputs.merge(
            frame, on="municipality_japanese", how="inner", validate="one_to_one"
        )
    audit_inputs = merged[["municipality_japanese", "municipality_english"]].merge(
        audit_inputs, on="municipality_japanese", how="left", validate="one_to_one"
    )
    if audit_inputs.isna().any().any():
        raise ValueError("Sex-specific certification inputs are incomplete")
    for column in LEVEL_COLUMNS:
        if not (
            audit_inputs[f"Total {column}"]
            == audit_inputs[f"Male {column}"] + audit_inputs[f"Female {column}"]
        ).all():
            raise ValueError(f"Certification sex reconciliation failed: {column}")
    return merged, audit_inputs


def prepare_main_table(merged: pd.DataFrame) -> pd.DataFrame:
    table = pd.DataFrame()
    table["Municipality"] = merged["Municipality"].astype(str)
    table["Municipality Observation Time"] = pd.to_datetime(
        merged["Municipality Observation Time"]
    ).dt.strftime("%Y-%m-%d %H:%M JST")
    table["Open Shelter Count"] = merged["Open Shelter Count"].astype(int)
    table["Evacuees"] = merged["Evacuees"].astype(int)
    for source, target in (
        ("shelter_roster_status", "Shelter Roster Availability"),
        ("shelter_occupancy_status", "Shelter Occupancy Availability"),
        ("water_status_availability", "Water Status Availability"),
        ("deployment_records_availability", "Deployment Records Availability"),
    ):
        mapped = merged[source].map(STATUS_MAPS[source])
        if mapped.isna().any():
            raise ValueError(f"Unmapped audit status in {source}")
        table[target] = mapped
    table["Residential Sex Input Status"] = "Available: 2020 population mesh"
    table["Functional Support Input Status"] = (
        "Available: Apr 2026 LTC by sex/care level"
    )
    table["Certification Sex Reconciliation Status"] = "Exact: male + female = total"
    table["Municipality Evidence Tier"] = merged["evidence_tier"].map(
        STATUS_MAPS["evidence_tier"]
    )
    table["Permitted Analysis Level"] = merged["permitted_analysis_level"].map(
        STATUS_MAPS["permitted_analysis_level"]
    )
    tier_two = merged["evidence_tier"].eq("tier_2_shelter_detailed")
    table["Reference-Period Alignment"] = (
        "Mixed: population 2020; LTC Apr 2026; municipality 14:00"
    )
    table.loc[tier_two, "Reference-Period Alignment"] = (
        "Mixed: population 2020; LTC Apr 2026; shelter 12:00; municipality 14:00"
    )
    table["Decisive Limitation"] = table["Municipality"].map(LIMITATIONS).fillna(
        DEFAULT_LIMITATION
    )
    if table[list(HEADERS)].isna().any().any():
        raise ValueError("Prepared evidence audit contains missing values")
    return table[list(HEADERS)]


def write_workbook(table: pd.DataFrame, audit_inputs: pd.DataFrame) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Coverage Audit"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "E2"
    sheet.sheet_view.zoomScale = 70
    sheet.append(list(HEADERS))

    for _, row in table.iterrows():
        sheet.append(row.tolist())

    last_row = sheet.max_row
    table_range = f"A1:O{last_row}"
    excel_table = Table(displayName="MunicipalityCoverageAudit", ref=table_range)
    excel_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(excel_table)

    header_fill = PatternFill("solid", fgColor="27445C")
    header_font = Font(name="Aptos", size=9, bold=True, color="FFFFFF")
    body_font = Font(name="Aptos", size=8, color="263746")
    thin_gray = Side(style="thin", color="D8DEE3")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="9EB1C0"))
    sheet.row_dimensions[1].height = 72

    for row in sheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=15):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
        row[1].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        row[2].alignment = Alignment(horizontal="right", vertical="center")
        row[3].alignment = Alignment(horizontal="right", vertical="center")
        sheet.row_dimensions[row[0].row].height = 58
    for column in ("C", "D"):
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0"

    available_fill = PatternFill("solid", fgColor="DCEDE5")
    partial_fill = PatternFill("solid", fgColor="FFF0CC")
    missing_fill = PatternFill("solid", fgColor="F5D7D7")
    for column in ("E", "F", "G", "H", "I", "J", "K"):
        target_range = f"{column}2:{column}{last_row}"
        sheet.conditional_formatting.add(
            target_range,
            FormulaRule(
                formula=[f'LEFT({column}2,9)="Available"'], fill=available_fill
            ),
        )
        sheet.conditional_formatting.add(
            target_range,
            FormulaRule(
                formula=[
                    f'OR(LEFT({column}2,7)="Partial",LEFT({column}2,8)="Identity",LEFT({column}2,10)="Identities")'
                ],
                fill=partial_fill,
            ),
        )
        sheet.conditional_formatting.add(
            target_range,
            FormulaRule(
                formula=[
                    f'OR(LEFT({column}2,9)="Not found",LEFT({column}2,14)="Available only")'
                ],
                fill=missing_fill,
            ),
        )
    sheet.conditional_formatting.add(
        f"K2:K{last_row}",
        FormulaRule(formula=['LEFT($K2,5)="Exact"'], fill=available_fill),
    )
    sheet.conditional_formatting.add(
        f"L2:L{last_row}",
        FormulaRule(formula=['LEFT($L2,6)="Tier 2"'], fill=available_fill),
    )
    sheet.conditional_formatting.add(
        f"L2:L{last_row}",
        FormulaRule(formula=['LEFT($L2,6)="Tier 1"'], fill=partial_fill),
    )

    widths = {
        "A": 16,
        "B": 18,
        "C": 10,
        "D": 10,
        "E": 20,
        "F": 20,
        "G": 19,
        "H": 22,
        "I": 20,
        "J": 22,
        "K": 21,
        "L": 19,
        "M": 24,
        "N": 25,
        "O": 37,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet["K1"].comment = Comment(
        "Validated against Audit Inputs. Exact requires male plus female to equal "
        "the published total for Support Levels 1-2, Care Levels 1-5, and the "
        "Certified Total in every municipality.",
        "Mike Li",
    )
    sheet.auto_filter.ref = table_range
    sheet.print_area = table_range
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.18
    sheet.page_margins.right = 0.18
    sheet.page_margins.top = 0.3
    sheet.page_margins.bottom = 0.3
    sheet.oddFooter.center.text = "Evidence availability—not evidence of service adequacy"
    sheet.oddFooter.center.size = 8
    sheet.oddFooter.center.color = "66737D"

    inputs = workbook.create_sheet("Audit Inputs")
    inputs.sheet_view.showGridLines = False
    input_headers = ["Municipality"]
    for sex in ("Total", "Male", "Female"):
        input_headers.extend(f"{sex} {column}" for column in LEVEL_COLUMNS)
    input_headers.append("Sex Reconciliation Status")
    inputs.append(input_headers)
    for row_number, (_, row) in enumerate(audit_inputs.iterrows(), start=2):
        input_values = [row["municipality_english"]]
        for sex in ("Total", "Male", "Female"):
            input_values.extend(int(row[f"{sex} {column}"]) for column in LEVEL_COLUMNS)
        input_values.append(
            f'=IF(SUMPRODUCT(--(B{row_number}:I{row_number}='
            f'(J{row_number}:Q{row_number}+R{row_number}:Y{row_number})))=8,'
            '"Exact: male + female = total","Mismatch")'
        )
        inputs.append(input_values)
    for cell in inputs[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    inputs.row_dimensions[1].height = 58
    for row in inputs.iter_rows(min_row=2, max_row=inputs.max_row):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.border = Border(bottom=thin_gray)
        row[0].alignment = Alignment(horizontal="left", vertical="center")
        row[-1].alignment = Alignment(horizontal="left", vertical="center")
        inputs.row_dimensions[row[0].row].height = 24
    inputs.column_dimensions["A"].width = 18
    for column in range(2, 26):
        inputs.column_dimensions[inputs.cell(row=1, column=column).column_letter].width = 12
    inputs.column_dimensions["Z"].width = 28
    inputs.freeze_panes = "B2"
    inputs.sheet_view.zoomScale = 65
    inputs.print_area = f"A1:Z{inputs.max_row}"
    inputs.page_setup.orientation = "landscape"
    inputs.page_setup.paperSize = inputs.PAPERSIZE_A3
    inputs.page_setup.fitToWidth = 1
    inputs.page_setup.fitToHeight = 1
    inputs.sheet_properties.pageSetUpPr.fitToPage = True

    notes = workbook.create_sheet("Notes")
    notes.sheet_view.showGridLines = False
    notes["A1"] = "Evidence audit notes and sources"
    notes["A1"].font = Font(name="Aptos Display", size=15, bold=True, color="27445C")
    notes.append([])
    notes.append(["Item", "Definition or boundary", "Source URL"])
    note_rows = (
        (
            "Municipality totals",
            "Observed at the common 12 August 2026 14:00 JST prefecture snapshot; all 11 municipalities reproduce 81 open shelters and 3,585 evacuees.",
            "https://www.pref.kumamoto.jp/uploaded/life/276971_869663_misc.pdf",
        ),
        (
            "Shelter evidence",
            "Kumamoto City and Yatsushiro City shelter tables are near-time 12:00 snapshots; availability does not establish fixed-stall inventory or functional operability.",
            "https://www.city.kumamoto.jp/kiji00372080/3_72080_514016_up_mh1811lq.pdf",
        ),
        (
            "Yatsushiro deployment evidence",
            "Reported placement and shelter records support conditional deployment scenarios; reported units are not automatically verified functional stalls.",
            "https://www.city.yatsushiro.lg.jp/kiji00326750/3_26750_158214_up_bjxkz7dq.pdf",
        ),
        (
            "Residential sex input",
            "The 2020 population mesh describes residents, not observed 2026 shelter users.",
            "https://www.e-stat.go.jp/gis/statmap-search?page=1&type=1&toukeiCode=00200521",
        ),
        (
            "Functional-support input",
            "April 2026 long-term-care certification by sex and care level is an administrative planning proxy, not disability prevalence or direct toilet-assistance need.",
            "https://www.mhlw.go.jp/topics/kaigo/osirase/jigyo/m26/xls/2604-h2-1.xlsx",
        ),
        (
            "Interpretation boundary",
            "Coverage status defines what can be analyzed; it does not show that sanitation service is sufficient, accessible, safe, gender-responsive, or functional.",
            "",
        ),
    )
    for values in note_rows:
        notes.append(values)
    for cell in notes[3]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
    notes.row_dimensions[3].height = 24
    for row in notes.iter_rows(min_row=4, max_row=notes.max_row, min_col=1, max_col=3):
        row[0].font = Font(name="Aptos", size=9, bold=True, color="263746")
        row[1].font = body_font
        row[2].font = Font(name="Aptos", size=8, color="0563C1", underline="single")
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
        notes.row_dimensions[row[0].row].height = 42
    notes.column_dimensions["A"].width = 27
    notes.column_dimensions["B"].width = 91
    notes.column_dimensions["C"].width = 78
    notes.freeze_panes = "A4"
    notes.sheet_view.zoomScale = 85
    notes.print_area = f"A1:C{notes.max_row}"
    notes.page_setup.orientation = "landscape"
    notes.page_setup.paperSize = notes.PAPERSIZE_A4
    notes.page_setup.fitToWidth = 1
    notes.page_setup.fitToHeight = 1
    notes.sheet_properties.pageSetUpPr.fitToPage = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)


def verify_workbook(expected: pd.DataFrame) -> None:
    workbook = load_workbook(OUTPUT, data_only=False, read_only=False)
    if workbook.sheetnames != ["Coverage Audit", "Audit Inputs", "Notes"]:
        raise ValueError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    sheet = workbook["Coverage Audit"]
    if sheet.max_row != 12 or sheet.max_column != 15:
        raise ValueError(f"Unexpected main-table dimensions: {sheet.max_row} x {sheet.max_column}")
    if tuple(cell.value for cell in sheet[1]) != HEADERS:
        raise ValueError("Main-table headers do not match the planned 15-column audit")
    if any(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for row in sheet.iter_rows()
        for cell in row
    ):
        raise ValueError("The article-facing Coverage Audit must contain static values")
    if [sheet.cell(row=row, column=11).value for row in range(2, 13)] != expected[
        "Certification Sex Reconciliation Status"
    ].tolist():
        raise ValueError("Certification reconciliation values changed in the main table")
    if sum(sheet.cell(row=row, column=3).value for row in range(2, 13)) != 81:
        raise ValueError("Workbook does not reproduce 81 open shelters")
    if sum(sheet.cell(row=row, column=4).value for row in range(2, 13)) != 3585:
        raise ValueError("Workbook does not reproduce 3,585 evacuees")
    inputs = workbook["Audit Inputs"]
    if inputs.max_row != 12 or inputs.max_column != 26:
        raise ValueError("Unexpected certification audit-input dimensions")
    if not all(str(inputs.cell(row=row, column=26).value).startswith("=IF(SUMPRODUCT") for row in range(2, 13)):
        raise ValueError("Certification reconciliation formulas are missing")
    if workbook["Notes"]["B9"].value is None:
        raise ValueError("Interpretation boundary is missing from Notes")
    if len(expected) != 11:
        raise ValueError("Validated table row count changed unexpectedly")
    workbook.close()


def main() -> None:
    merged, audit_inputs = read_and_validate_inputs()
    table = prepare_main_table(merged)
    write_workbook(table, audit_inputs)
    verify_workbook(table)
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print("Audit: 11 municipalities, 15 main columns, 81 shelters, 3,585 evacuees")
    print("Certification reconciliation: exact across 8 published measures")


if __name__ == "__main__":
    main()
