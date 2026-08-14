#!/usr/bin/env python3
"""Toilet Rebalancing Scenario Performance and Robustness.

Plan: Summarize three predeclared temporary-toilet mobility scenarios in a
compact decision table, with the complete metric record retained separately.
Framework: Section 6 donor capacity, integer transfer, residual shortfall, and
resolved-site coverage definitions. Results are conditional planning screens,
not verified dispatch instructions or functional-capacity findings.
"""

from __future__ import annotations

from pathlib import Path
import re

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

import figure_toilet_rebalancing_performance_under_resource_mobility_scenarios as scenarios


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/processed/shelter_equity_scenarios_preprocessed.parquet"
OUTPUT = (
    ROOT
    / "data/results/tables/"
    "Table_toilet_rebalancing_scenario_performance_and_robustness.xlsx"
)

JAPANESE_PATTERN = re.compile(r"[\u3040-\u30ff\u3400-\u9fff]")
SCENARIO_NAMES = {
    "none": "No rebalancing",
    "zero": "Zero-occupancy donor only",
    "full": "Full reported-surplus mobility",
}
MOBILITY_RULES = {
    "none": "No reported units movable",
    "zero": "Only reported units at zero-occupancy shelters movable",
    "full": "All reported units above each site's prolonged requirement movable",
}
BOUNDARIES = {
    "none": "Observed temporary-toilet-only screening baseline",
    "zero": "Conservative immediate-transfer opportunity, conditional on field checks",
    "full": "Best-case spatial-rebalancing bound; mobility-sensitive",
}
FULL_HEADERS = (
    "Scenario",
    "Mobility Rule",
    "Eligible Donor Units",
    "Transferred Units",
    "Recipient Shelters Served",
    "Residual Screening Shortfall",
    "Shelters with Residual Shortfall",
    "Resolved-Site Demand Coverage (Evacuees / Functional Support / Female Functional Support)",
    "Unresolved-Site Demand (Evacuees / Functional Support / Female Functional Support)",
    "Mean Transfer Distance (km)",
    "Maximum Transfer Distance (km)",
    "Interpretation Boundary",
)
MAIN_HEADERS = (
    "Metric",
    "No Rebalancing",
    "Zero-Occupancy Donor Only",
    "Full Reported-Surplus Mobility",
)


def evaluate_detail(screen: pd.DataFrame, scenario: str) -> dict[str, object]:
    capacities = scenarios.donor_capacity(screen, scenario)
    eligible_units = int(sum(capacities.values()))
    residual = screen["Screening Shortfall"].astype(int).to_dict()
    recipients = screen.loc[screen["Screening Shortfall"].gt(0)].sort_values(
        [
            "Verification Tier",
            "Screening Shortfall",
            "Estimated Functional Support Evacuees",
            "Estimated Female Functional Support Evacuees",
            "Evacuees",
            "Shelter Number",
        ],
        ascending=[True, False, False, False, False, True],
        kind="stable",
    )
    flows: list[tuple[int, int, int, float]] = []
    for recipient_index, recipient in recipients.iterrows():
        while residual[int(recipient_index)] > 0:
            eligible = [
                donor_index
                for donor_index, capacity in capacities.items()
                if capacity > 0 and donor_index != recipient_index
            ]
            if not eligible:
                break
            donor_index = min(
                eligible,
                key=lambda index: (
                    scenarios.haversine_km(screen.loc[index], recipient),
                    int(screen.loc[index, "Shelter Number"]),
                ),
            )
            quantity = min(capacities[donor_index], residual[int(recipient_index)])
            distance = scenarios.haversine_km(screen.loc[donor_index], recipient)
            capacities[donor_index] -= quantity
            residual[int(recipient_index)] -= quantity
            flows.append((donor_index, int(recipient_index), int(quantity), float(distance)))

    transferred = int(sum(quantity for _, _, quantity, _ in flows))
    served = sorted({recipient for _, recipient, _, _ in flows})
    unresolved = [int(index) for index in recipients.index if residual[int(index)] > 0]
    resolved = [int(index) for index in recipients.index if residual[int(index)] == 0]
    columns = (
        "Evacuees",
        "Estimated Functional Support Evacuees",
        "Estimated Female Functional Support Evacuees",
    )

    def demand(indices: list[int], column: str) -> float:
        return float(screen.loc[indices, column].sum()) if indices else 0.0

    all_recipients = list(map(int, recipients.index))
    coverage = []
    unresolved_demand = []
    for column in columns:
        total = demand(all_recipients, column)
        resolved_total = demand(resolved, column)
        coverage.append(100.0 * resolved_total / total if total else 0.0)
        unresolved_demand.append(demand(unresolved, column))
    mean_distance = (
        sum(quantity * distance for _, _, quantity, distance in flows) / transferred
        if transferred
        else 0.0
    )
    max_distance = max((distance for _, _, _, distance in flows), default=0.0)
    return {
        "Scenario": SCENARIO_NAMES[scenario],
        "Mobility Rule": MOBILITY_RULES[scenario],
        "Eligible Donor Units": eligible_units,
        "Transferred Units": transferred,
        "Recipient Shelters Served": len(served),
        "Residual Screening Shortfall": int(sum(residual.values())),
        "Shelters with Residual Shortfall": len(unresolved),
        "Resolved-Site Demand Coverage (Evacuees / Functional Support / Female Functional Support)": (
            f"{coverage[0]:.1f}% / {coverage[1]:.1f}% / {coverage[2]:.1f}%"
        ),
        "Unresolved-Site Demand (Evacuees / Functional Support / Female Functional Support)": (
            f"{unresolved_demand[0]:.0f} / {unresolved_demand[1]:.1f} / {unresolved_demand[2]:.1f}"
        ),
        "Mean Transfer Distance (km)": mean_distance,
        "Maximum Transfer Distance (km)": max_distance,
        "Interpretation Boundary": BOUNDARIES[scenario],
    }


def build_table(frame: pd.DataFrame) -> pd.DataFrame:
    base = frame.loc[
        frame["Municipality"].eq("八代市")
        & frame["Scenario"].astype("string").eq("base")
    ].copy().reset_index(drop=True)
    screen = scenarios.construct_screen(base)
    table = pd.DataFrame(
        [evaluate_detail(screen, scenario) for scenario in scenarios.SCENARIOS]
    )
    expected = {
        "No rebalancing": (0, 0, 0, 31, 11),
        "Zero-occupancy donor only": (5, 5, 3, 26, 9),
        "Full reported-surplus mobility": (72, 31, 11, 0, 0),
    }
    for _, row in table.iterrows():
        observed = (
            int(row["Eligible Donor Units"]),
            int(row["Transferred Units"]),
            int(row["Recipient Shelters Served"]),
            int(row["Residual Screening Shortfall"]),
            int(row["Shelters with Residual Shortfall"]),
        )
        if observed != expected[str(row["Scenario"])]:
            raise ValueError(f"Unexpected scenario result: {row['Scenario']} {observed}")
    if table[list(FULL_HEADERS)].isna().any().any():
        raise ValueError("Scenario table contains missing values")
    return table[list(FULL_HEADERS)]


def add_table_style(sheet, name: str, reference: str) -> None:
    excel_table = Table(displayName=name, ref=reference)
    excel_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(excel_table)


def write_workbook(table: pd.DataFrame) -> None:
    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    sheet = workbook.active
    sheet.title = "Scenario Performance"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "B2"
    sheet.sheet_view.zoomScale = 90
    sheet.append(list(MAIN_HEADERS))
    scenario_rows = [row for _, row in table.iterrows()]
    metric_values = (
        (
            "Mobility Rule",
            [row["Mobility Rule"] for row in scenario_rows],
        ),
        (
            "Resource Movement\n(Eligible / Transferred Units)",
            [
                f"{int(row['Eligible Donor Units'])} / {int(row['Transferred Units'])}"
                for row in scenario_rows
            ],
        ),
        (
            "Recipient Outcomes\n(Received Units / Still Short)",
            [
                f"{int(row['Recipient Shelters Served'])} / {int(row['Shelters with Residual Shortfall'])}"
                for row in scenario_rows
            ],
        ),
        (
            "Residual Shortfall (Units)",
            [int(row["Residual Screening Shortfall"]) for row in scenario_rows],
        ),
        (
            "Resolved-Site Demand Coverage\n(Evacuees / Functional Support / Female Functional Support)",
            [
                row[
                    "Resolved-Site Demand Coverage (Evacuees / Functional Support / Female Functional Support)"
                ]
                for row in scenario_rows
            ],
        ),
        (
            "Unresolved-Site Demand\n(Evacuees / Functional Support / Female Functional Support)",
            [
                row[
                    "Unresolved-Site Demand (Evacuees / Functional Support / Female Functional Support)"
                ]
                for row in scenario_rows
            ],
        ),
        (
            "Transfer Distance\n(Mean / Maximum km)",
            [
                f"{float(row['Mean Transfer Distance (km)']):.2f} / {float(row['Maximum Transfer Distance (km)']):.2f}"
                for row in scenario_rows
            ],
        ),
        (
            "Interpretation Boundary",
            [row["Interpretation Boundary"] for row in scenario_rows],
        ),
    )
    for metric, values in metric_values:
        sheet.append([metric, *values])
    add_table_style(sheet, "CompactScenarioPerformance", "A1:D9")

    navy = "27445C"
    text = "263746"
    border = Side(style="thin", color="D8DEE3")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=8.6, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="9EB1C0"))
    sheet.row_dimensions[1].height = 52
    scenario_colors = {2: "E8EDF1", 3: "F4DEB6", 4: "DCEDE5"}
    for row_number in range(2, 10):
        for cell in sheet[row_number]:
            cell.font = Font(name="Aptos", size=8.2, color=text)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=border, left=border)
        sheet.cell(row_number, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.cell(row_number, 1).fill = PatternFill("solid", fgColor="37566F")
        sheet.cell(row_number, 1).font = Font(name="Aptos", size=8.2, bold=True, color="FFFFFF")
        for column_number, fill_color in scenario_colors.items():
            sheet.cell(row_number, column_number).fill = PatternFill("solid", fgColor=fill_color)
        if row_number in (2, 9):
            for column_number in range(2, 5):
                sheet.cell(row_number, column_number).alignment = Alignment(
                    horizontal="left", vertical="center", wrap_text=True
                )
        sheet.row_dimensions[row_number].height = 50
    for column_number in (2, 3):
        sheet.cell(5, column_number).fill = PatternFill("solid", fgColor="F5D7D7")
    sheet.cell(5, 4).fill = PatternFill("solid", fgColor="CDE2D7")
    sheet.row_dimensions[2].height = 62
    sheet.row_dimensions[6].height = 58
    sheet.row_dimensions[7].height = 58
    sheet.row_dimensions[9].height = 64
    for column, width in {
        "A": 43,
        "B": 34,
        "C": 41,
        "D": 44,
    }.items():
        sheet.column_dimensions[column].width = width
    sheet["A5"].comment = Comment(
        "Residual temporary-toilet-only screening shortfall excludes fixed stalls and toilet-car stall equivalents.",
        "OpenAI",
    )
    sheet["A6"].comment = Comment(
        "Coverage is the share of demand at the initial eleven shortfall shelters that becomes fully resolved in each scenario.",
        "OpenAI",
    )
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = "A1:D9"

    full = workbook.create_sheet("Full Scenario Record")
    full.sheet_view.showGridLines = False
    full.append(list(FULL_HEADERS))
    for _, row in table.iterrows():
        full.append(row.tolist())
    add_table_style(full, "FullScenarioPerformance", "A1:L4")
    for cell in full[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=8.2, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    full.row_dimensions[1].height = 68
    for row in full.iter_rows(min_row=2, max_row=4):
        for cell in row:
            cell.font = Font(name="Aptos", size=8, color=text)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=border)
        row[0].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        row[1].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        row[11].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        full.row_dimensions[row[0].row].height = 58
    for column, width in {
        "A": 28, "B": 48, "C": 14, "D": 14, "E": 16, "F": 16,
        "G": 17, "H": 34, "I": 34, "J": 16, "K": 18, "L": 42,
    }.items():
        full.column_dimensions[column].width = width
    for row_number in range(2, 5):
        full.cell(row_number, 10).number_format = "0.00"
        full.cell(row_number, 11).number_format = "0.00"
    full.freeze_panes = "C2"
    full.auto_filter.ref = "A1:L4"

    notes = workbook.create_sheet("Notes")
    notes.sheet_view.showGridLines = False
    notes.column_dimensions["A"].width = 32
    notes.column_dimensions["B"].width = 108
    notes.append(["Field", "Definition"])
    note_rows = [
        ("Study set", "Thirty-eight Yatsushiro shelters with complete occupancy, reported temporary-toilet, toilet-car, water-status, and coordinate records at the matched observation."),
        ("Initial shortfall set", "Eleven shelters with 31 units of positive prolonged temporary-toilet-only screening shortfall."),
        ("Priority rule", "Recipients are processed by verification tier, screening shortfall, equity-sensitive demand, evacuees, and stable shelter ID; the nearest eligible donor is selected at each step."),
        ("Resolved-site coverage", "Share of evacuee or synthetic equity demand at the initial shortfall shelters whose screening shortfall becomes zero."),
        ("Distance", "Unit-weighted mean and maximum great-circle distances. District-anchor coordinates are not dispatch-ready routes."),
        ("Evidence boundary", "Reported units are assumed serviceable and movable within each scenario. Fixed stalls, toilet-car stall equivalents, ownership, road access, setup, wastewater, collection, and field operability are not verified."),
        ("Robustness conclusion", "The conservative zero-occupancy scenario improves but does not resolve the screen; complete resolution occurs only under full reported-surplus mobility, so the mitigation result is mobility-sensitive."),
    ]
    for row in note_rows:
        notes.append(row)
    for cell in notes[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
    for row in notes.iter_rows(min_row=2, max_row=notes.max_row):
        row[0].font = Font(name="Aptos", size=9, bold=True, color=text)
        row[1].font = Font(name="Aptos", size=9, color=text)
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = Border(bottom=border)
        notes.row_dimensions[row[0].row].height = 50
    notes.freeze_panes = "A2"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)


def validate_output(table: pd.DataFrame) -> None:
    workbook = load_workbook(OUTPUT, data_only=False)
    if workbook.sheetnames != ["Scenario Performance", "Full Scenario Record", "Notes"]:
        raise ValueError("Unexpected workbook sheet structure")
    if workbook["Scenario Performance"].max_row != 9 or workbook["Scenario Performance"].max_column != 4:
        raise ValueError("Unexpected compact-table dimensions")
    if workbook["Full Scenario Record"].max_row != 4 or workbook["Full Scenario Record"].max_column != 12:
        raise ValueError("Unexpected full-record dimensions")
    formulas = [
        cell
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    if len(formulas) != 0:
        raise ValueError(f"Unexpected formula count: {len(formulas)}")
    japanese_cells = [
        f"{worksheet.title}!{cell.coordinate}"
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and JAPANESE_PATTERN.search(cell.value)
    ]
    if japanese_cells:
        raise ValueError(f"Japanese text remains in workbook: {japanese_cells[:5]}")
    if int(table["Transferred Units"].sum()) != 36:
        raise ValueError("Unexpected transferred-unit total across sensitivity scenarios")


def main() -> None:
    table = build_table(pd.read_parquet(SOURCE))
    write_workbook(table)
    validate_output(table)
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print("Main table: 8 metric rows x 4 columns; full 12-column record retained")
    print("Transferred units: 0 / 5 / 31")
    print("Residual screening shortfall: 31 / 26 / 0")


if __name__ == "__main__":
    main()
