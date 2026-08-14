#!/usr/bin/env python3
"""Priority Shelter Gender and Accessible Service Packages.

Plan: Convert the eleven positive-screening-shortfall Yatsushiro shelters into
transparent women, men, accessible-parity, containment, hygiene, safety,
menstrual-waste, and accessible-route action packages.
Framework: Section 5 equity-sensitive conditional mitigation; Section 6
general-unit addition, women-to-men designation, and separate accessible parity
screen; Section 7 service-package verification workflow. Actions are planning
recommendations, not observed facility features or compliance findings.
"""

from __future__ import annotations

from pathlib import Path
import re

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

import figure_gender_and_functional_support_toilet_service_packages as packages
from table_shelter_level_equity_demand_scenario_estimates import (
    AREA_LABELS,
    ENGLISH_SHELTER_NAMES,
)


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/processed/shelter_equity_scenarios_preprocessed.parquet"
OUTPUT = (
    ROOT
    / "data/results/tables/"
    "Table_priority_shelter_gender_and_accessible_service_packages.xlsx"
)

JAPANESE_PATTERN = re.compile(r"[\u3040-\u30ff\u3400-\u9fff]")
WATER_LABELS = {"〇": "Available", "○": "Available", "×": "Unavailable"}
HEADERS = (
    "Shelter (ID / Name / Area)",
    "Evacuees",
    "Estimated Female Evacuees",
    "Estimated Functional-Support Evacuees",
    "Estimated Female Functional-Support Evacuees",
    "Proposed General-Unit Addition",
    "Women Planning Designation",
    "Men Planning Designation",
    "Accessible Unit Parity Screen",
    "Water Status",
    "Toilet Cars",
    "Containment or Collection Action",
    "Lighting and Locking Action",
    "Handwashing Action",
    "Menstrual-Waste Action",
    "Accessible-Route Action",
)
MAIN_HEADERS = (
    "Shelter (ID / Name / Area)",
    "Evacuees",
    "Estimated Demand (Women / Functional Support / Women + Functional Support)",
    "Proposed Units (Total / Women / Men)",
    "Accessible Unit Parity Screen",
    "Operational Context (Water / Toilet Cars)",
    "Containment or Collection",
    "Safety and Hygiene (Lighting + Locks / Handwashing)",
    "Menstrual-Waste Support",
    "Accessible Route",
)


def build_package(frame: pd.DataFrame) -> pd.DataFrame:
    base = frame.loc[
        frame["Municipality"].eq("八代市")
        & frame["Scenario"].astype("string").eq("base")
    ].copy()
    priority = packages.construct_priority_frame(base)
    packages.validate(base, priority)
    priority["Stable Shelter ID"] = (
        "Y" + priority["Shelter Number"].astype(int).astype(str).str.zfill(2)
    )
    table = pd.DataFrame()
    table["Shelter (ID / Name / Area)"] = [
        f"{shelter_id} — {ENGLISH_SHELTER_NAMES[shelter_id]}\n{AREA_LABELS[str(area)]}"
        for shelter_id, area in zip(
            priority["Stable Shelter ID"], priority["District"], strict=True
        )
    ]
    table["Evacuees"] = priority["Evacuees"].astype(int)
    table["Estimated Female Evacuees"] = priority["Estimated Female Evacuees"].astype(float)
    table["Estimated Functional-Support Evacuees"] = priority[
        "Estimated Functional Support Evacuees"
    ].astype(float)
    table["Estimated Female Functional-Support Evacuees"] = priority[
        "Estimated Female Functional Support Evacuees"
    ].astype(float)
    table["Proposed General-Unit Addition"] = priority["General Unit Addition"].astype(int)
    table["Women Planning Designation"] = priority["Women Designation"].astype(int)
    table["Men Planning Designation"] = priority["Men Designation"].astype(int)
    table["Accessible Unit Parity Screen"] = priority[
        "Prolonged Accessible Unit Parity Screen"
    ].astype(int)
    table["Water Status"] = priority["Water Status"].astype(str).map(WATER_LABELS)
    table["Toilet Cars"] = priority["Toilet Cars"].astype(int)
    unavailable = table["Water Status"].eq("Unavailable")
    table["Containment or Collection Action"] = "Confirm backup collection"
    table.loc[unavailable, "Containment or Collection Action"] = "Set storage + collection"
    table["Lighting and Locking Action"] = "Verify / provide"
    table["Handwashing Action"] = "Provide"
    table["Menstrual-Waste Action"] = "Provide bin + supplies"
    table["Accessible-Route Action"] = "Verify / clear"
    if table[list(HEADERS)].isna().any().any() or len(table) != 11:
        raise ValueError("Priority service-package table is incomplete")
    if int(table["Proposed General-Unit Addition"].sum()) != 31:
        raise ValueError("Expected 31 general-unit additions")
    if int(table["Women Planning Designation"].sum()) != 26:
        raise ValueError("Expected 26 women-designated units")
    if int(table["Men Planning Designation"].sum()) != 5:
        raise ValueError("Expected 5 men-designated units")
    if int(table["Accessible Unit Parity Screen"].sum()) != 12:
        raise ValueError("Expected accessible-parity screen total of 12")
    return table[list(HEADERS)]


def write_workbook(table: pd.DataFrame) -> None:
    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    sheet = workbook.active
    sheet.title = "Service Packages"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "B2"
    sheet.sheet_view.zoomScale = 88
    sheet.append(list(MAIN_HEADERS))
    for row_number, (_, row) in enumerate(table.iterrows(), start=2):
        sheet.append(
            [
                row.iloc[0],
                f"='Full Service Package Record'!B{row_number}",
                (
                    f'=TEXT(\'Full Service Package Record\'!C{row_number},"0.0")&" / "&'
                    f'TEXT(\'Full Service Package Record\'!D{row_number},"0.0")&" / "&'
                    f'TEXT(\'Full Service Package Record\'!E{row_number},"0.0")'
                ),
                (
                    f'=TEXT(\'Full Service Package Record\'!F{row_number},"0")&" / "&'
                    f'TEXT(\'Full Service Package Record\'!G{row_number},"0")&" / "&'
                    f'TEXT(\'Full Service Package Record\'!H{row_number},"0")'
                ),
                f"='Full Service Package Record'!I{row_number}",
                (
                    f'=\'Full Service Package Record\'!J{row_number}&" / "&'
                    f'TEXT(\'Full Service Package Record\'!K{row_number},"0")'
                ),
                f"='Full Service Package Record'!L{row_number}",
                (
                    f'=\'Full Service Package Record\'!M{row_number}&" / "&'
                    f"'Full Service Package Record'!N{row_number}"
                ),
                f"='Full Service Package Record'!O{row_number}",
                f"='Full Service Package Record'!P{row_number}",
            ]
        )

    last_row = sheet.max_row
    excel_table = Table(displayName="CompactPriorityServicePackages", ref=f"A1:J{last_row}")
    excel_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(excel_table)
    navy = "27445C"
    text = "263746"
    border = Side(style="thin", color="D8DEE3")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=8.2, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="9EB1C0"))
    sheet.row_dimensions[1].height = 72
    for row in sheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=10):
        for column_index, cell in enumerate(row):
            cell.font = Font(name="Aptos", size=7.8, color=text)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(
                left=border if column_index > 0 else Side(style=None), bottom=border
            )
        row[1].alignment = Alignment(horizontal="right", vertical="center")
        row[4].alignment = Alignment(horizontal="right", vertical="center")
        for index in (2, 3, 5, 6, 7, 8, 9):
            row[index].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 48
    for cell in sheet["B"][1:]:
        cell.number_format = "#,##0"
    for cell in sheet["E"][1:]:
        cell.number_format = "#,##0"

    quantity_fills = {"C": "E9CADC", "D": "F3CEC5", "E": "CDE2D7"}
    for column, color in quantity_fills.items():
        for cell in sheet[column][1:]:
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(name="Aptos", size=7.8, bold=True, color=text)
    action_fills = {"G": "F4DEB6", "H": "CDE2D7", "I": "E9CADC", "J": "CDE2D7"}
    for column, color in action_fills.items():
        for cell in sheet[column][1:]:
            cell.fill = PatternFill("solid", fgColor=color)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_number, status in enumerate(table["Water Status"], start=2):
        color = "F5D7D7" if status == "Unavailable" else "DCEDE5"
        sheet.cell(row_number, 6).fill = PatternFill("solid", fgColor=color)
        if status == "Unavailable":
            sheet.cell(row_number, 7).fill = PatternFill("solid", fgColor="F3CEC5")

    widths = {
        "A": 42,
        "B": 10,
        "C": 29,
        "D": 25,
        "E": 19,
        "F": 22,
        "G": 25,
        "H": 29,
        "I": 23,
        "J": 20,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet["D1"].comment = Comment(
        "General additions equal the positive prolonged requirement minus reported temporary toilets. Existing fixed stalls and toilet-car stall equivalents are not included.",
        "OpenAI",
    )
    sheet["E1"].comment = Comment(
        "Accessible parity is a separate Base-scenario screen and is not subtracted from or added to the women and men designations.",
        "OpenAI",
    )
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = f"A1:J{last_row}"

    inputs = workbook.create_sheet("Full Service Package Record")
    inputs.sheet_view.showGridLines = False
    inputs.append(list(HEADERS))
    for row_number, (_, row) in enumerate(table.iterrows(), start=2):
        values = row.tolist()
        values[6] = f"=ROUNDUP(0.75*F{row_number},0)"
        values[7] = f"=F{row_number}-G{row_number}"
        inputs.append(values)
    for cell in inputs[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=8.2, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    inputs.row_dimensions[1].height = 66
    for row in inputs.iter_rows(min_row=2, max_row=inputs.max_row):
        for cell in row:
            cell.font = Font(name="Aptos", size=7.8, color=text)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=border)
        inputs.row_dimensions[row[0].row].height = 44
    full_widths = {
        "A": 39,
        "B": 10,
        "C": 16,
        "D": 19,
        "E": 23,
        "F": 17,
        "G": 16,
        "H": 15,
        "I": 17,
        "J": 16,
        "K": 10,
        "L": 22,
        "M": 18,
        "N": 15,
        "O": 19,
        "P": 17,
    }
    for column, width in full_widths.items():
        inputs.column_dimensions[column].width = width
    inputs.freeze_panes = "B2"
    inputs.auto_filter.ref = f"A1:P{inputs.max_row}"
    full_table = Table(displayName="FullServicePackageRecord", ref=f"A1:P{inputs.max_row}")
    full_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    inputs.add_table(full_table)

    notes = workbook.create_sheet("Notes")
    notes.sheet_view.showGridLines = False
    notes.column_dimensions["A"].width = 31
    notes.column_dimensions["B"].width = 108
    notes.append(["Field", "Definition"])
    note_rows = [
        ("Priority set", "The eleven Yatsushiro shelters with a positive prolonged temporary-toilet-only screening shortfall."),
        ("Women and men designations", "Incremental general units are designated 3:1 for women and men using ceiling allocation for women. These designations do not reveal or correct the allocation of existing stalls."),
        ("Accessible parity", "A separate Base-scenario planning screen derived from estimated functional-support evacuees. It is not an observed accessible-toilet requirement or compliance finding."),
        ("Water and containment", "Unavailable water triggers non-flush storage and collection planning. Available water still requires wastewater and backup collection verification."),
        ("Service package", "Lighting, locking, handwashing, menstrual-waste support, and accessible routes are actions to verify or provide, not observed facility conditions."),
        ("Evidence boundary", "Female and functional-support quantities are synthetic expectations. The package is conditional until equipment, water, wastewater, collection, safety, hygiene, and routes are field verified."),
    ]
    for row in note_rows:
        notes.append(row)
    for cell in notes[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
    for row in notes.iter_rows(min_row=2, max_row=notes.max_row):
        row[0].font = Font(name="Aptos", size=9, bold=True, color=text)
        row[1].font = Font(name="Aptos", size=9, color=text)
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = Border(bottom=border)
        notes.row_dimensions[row[0].row].height = 48
    notes.freeze_panes = "A2"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)


def validate_output(table: pd.DataFrame) -> None:
    workbook = load_workbook(OUTPUT, data_only=False)
    if workbook.sheetnames != ["Service Packages", "Full Service Package Record", "Notes"]:
        raise ValueError("Unexpected workbook sheet structure")
    sheet = workbook["Service Packages"]
    if sheet.max_row != 12 or sheet.max_column != 10:
        raise ValueError("Unexpected main-table dimensions")
    formulas = [
        cell
        for workbook_sheet in workbook.worksheets
        for row in workbook_sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    if len(formulas) != 121:
        raise ValueError(f"Unexpected formula count: {len(formulas)}")
    japanese_cells = []
    for workbook_sheet in workbook.worksheets:
        for row in workbook_sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and JAPANESE_PATTERN.search(cell.value):
                    japanese_cells.append(f"{workbook_sheet.title}!{cell.coordinate}")
    if japanese_cells:
        raise ValueError(f"Japanese text remains in workbook: {japanese_cells[:5]}")
    if int(table["Women Planning Designation"].sum()) != 26:
        raise ValueError("Women designation total changed")
    if int(table["Men Planning Designation"].sum()) != 5:
        raise ValueError("Men designation total changed")


def main() -> None:
    table = build_package(pd.read_parquet(SOURCE))
    write_workbook(table)
    validate_output(table)
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print("Main table: 11 priority shelters x 10 columns; full 16-column record retained")
    print("General additions: 31; women designation: 26; men designation: 5")
    print("Base accessible-parity screen: 12 separate units")


if __name__ == "__main__":
    main()
