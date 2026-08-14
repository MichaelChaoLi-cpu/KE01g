#!/usr/bin/env python3
"""Robust Shelter Verification Priorities across Scenarios.

Plan: Identify the 15 shelters with the smallest worst cross-city verification
rank across the low, base, and high functional-support overlap scenarios while
retaining women, water, and reported deployment context.
Framework: Section 5 scenario-stability contrast; Section 6 global scenario
ranks omega and worst rank M; Section 7 verification-priority robustness
workflow. The result prioritizes field verification and is not a deployment
optimization or an observed functional-capacity gap.
"""

from __future__ import annotations

from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from table_shelter_level_equity_demand_scenario_estimates import (
    AREA_LABELS,
    ENGLISH_SHELTER_NAMES,
)


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/processed/shelter_equity_scenarios_preprocessed.parquet"
OUTPUT = (
    ROOT
    / "data/results/tables/"
    "Table_robust_shelter_verification_priorities_across_scenarios.xlsx"
)

SCENARIOS = ("low", "base", "high")
JAPANESE_PATTERN = re.compile(r"[\u3040-\u30ff\u3400-\u9fff]")
MUNICIPALITY_LABELS = {"八代市": "Yatsushiro City", "熊本市": "Kumamoto City"}
MUNICIPALITY_PREFIXES = {"八代市": "Y", "熊本市": "K"}
WATER_LABELS = {
    "〇": "Available",
    "○": "Available",
    "△": "Partially available",
    "×": "Unavailable",
}

HEADERS = (
    "Robust Priority Order",
    "Municipality",
    "Shelter Name",
    "District",
    "Ward",
    "Evacuees",
    "Estimated Female Evacuees",
    "Functional-Support Demand (Low / Base / High)",
    "Estimated Female Functional-Support Evacuees (Base)",
    "Water Status",
    "Temporary Toilets Installed",
    "Toilet Cars",
    "Cross-Scenario Ranks (Low / Base / High)",
    "Worst-Scenario Rank",
)


def stable_id(frame: pd.DataFrame) -> pd.Series:
    prefixes = frame["Municipality"].map(MUNICIPALITY_PREFIXES)
    if prefixes.isna().any():
        raise ValueError("Unexpected municipality in matched shelter sample")
    numbers = frame["Shelter Number"].astype(int).astype(str).str.zfill(2)
    return prefixes + numbers


def build_ranking(frame: pd.DataFrame) -> pd.DataFrame:
    required = {
        "Municipality",
        "Shelter Number",
        "Shelter Name",
        "District",
        "Ward",
        "Scenario",
        "Evacuees",
        "Estimated Female Evacuees",
        "Estimated Functional Support Evacuees",
        "Estimated Female Functional Support Evacuees",
        "Water Status",
        "Temporary Toilets Installed",
        "Toilet Cars",
        "Location Resolution",
    }
    missing = required.difference(frame.columns)
    if missing:
        raise KeyError(f"Missing required columns: {sorted(missing)}")
    frame = frame.loc[
        frame["Municipality"].isin(MUNICIPALITY_LABELS)
        & frame["Scenario"].astype("string").isin(SCENARIOS)
    ].copy()
    if len(frame) != 159:
        raise ValueError("Expected 53 shelters under each of three scenarios")
    if frame.duplicated(["Municipality", "Shelter Number", "Scenario"]).any():
        raise ValueError("Shelter-scenario keys must be unique")
    if set(frame["Scenario"].astype(str)) != set(SCENARIOS):
        raise ValueError("Expected exactly low, base, and high scenarios")
    if frame[
        [
            "Municipality",
            "Shelter Number",
            "Shelter Name",
            "Evacuees",
            "Estimated Female Evacuees",
            "Estimated Functional Support Evacuees",
            "Estimated Female Functional Support Evacuees",
        ]
    ].isna().any().any():
        raise ValueError("Ranking inputs must be complete")
    frame["Stable Shelter ID"] = stable_id(frame)

    ranked_frames: list[pd.DataFrame] = []
    for scenario in SCENARIOS:
        scenario_frame = frame.loc[frame["Scenario"].astype(str).eq(scenario)].copy()
        scenario_frame = scenario_frame.sort_values(
            [
                "Estimated Functional Support Evacuees",
                "Estimated Female Functional Support Evacuees",
                "Evacuees",
                "Stable Shelter ID",
            ],
            ascending=[False, False, False, True],
            kind="stable",
        )
        scenario_frame["Cross-City Verification Rank"] = np.arange(1, 54)
        ranked_frames.append(scenario_frame)
    ranked = pd.concat(ranked_frames, ignore_index=True)

    for scenario in SCENARIOS:
        values = ranked.loc[
            ranked["Scenario"].astype(str).eq(scenario),
            "Cross-City Verification Rank",
        ].astype(int)
        if sorted(values.tolist()) != list(range(1, 54)):
            raise ValueError(f"Invalid cross-city rank permutation: {scenario}")

    base_columns = [
        "Stable Shelter ID",
        "Municipality",
        "Shelter Number",
        "Shelter Name",
        "District",
        "Ward",
        "Evacuees",
        "Estimated Female Evacuees",
        "Estimated Female Functional Support Evacuees",
        "Water Status",
        "Temporary Toilets Installed",
        "Toilet Cars",
        "Location Resolution",
    ]
    base = ranked.loc[ranked["Scenario"].astype(str).eq("base"), base_columns].copy()
    base = base.rename(
        columns={
            "Estimated Female Functional Support Evacuees": (
                "Estimated Female Functional-Support Evacuees (Base)"
            )
        }
    )
    demand_wide = (
        ranked.pivot(
            index="Stable Shelter ID",
            columns="Scenario",
            values="Estimated Functional Support Evacuees",
        )
        .rename(columns={scenario: f"FS Demand {scenario.title()}" for scenario in SCENARIOS})
        .reset_index()
    )
    rank_wide = (
        ranked.pivot(
            index="Stable Shelter ID",
            columns="Scenario",
            values="Cross-City Verification Rank",
        )
        .rename(columns={scenario: f"Rank {scenario.title()}" for scenario in SCENARIOS})
        .reset_index()
    )
    wide = base.merge(demand_wide, on="Stable Shelter ID", validate="one_to_one")
    wide = wide.merge(rank_wide, on="Stable Shelter ID", validate="one_to_one")
    rank_columns = [f"Rank {scenario.title()}" for scenario in SCENARIOS]
    wide["Worst-Scenario Rank"] = wide[rank_columns].max(axis=1).astype(int)
    wide["Rank Span"] = (
        wide[rank_columns].max(axis=1) - wide[rank_columns].min(axis=1)
    ).astype(int)

    demand_columns = [f"FS Demand {scenario.title()}" for scenario in SCENARIOS]
    if not (
        wide["FS Demand Low"].le(wide["FS Demand Base"])
        & wide["FS Demand Base"].le(wide["FS Demand High"])
    ).all():
        raise ValueError("Functional-support demand is not monotonic across scenarios")
    if wide[demand_columns].lt(0).any().any():
        raise ValueError("Functional-support demand cannot be negative")

    wide = wide.sort_values(
        [
            "Worst-Scenario Rank",
            "FS Demand Base",
            "Estimated Female Functional-Support Evacuees (Base)",
            "Evacuees",
            "Stable Shelter ID",
        ],
        ascending=[True, False, False, False, True],
        kind="stable",
    ).reset_index(drop=True)
    wide["Robust Priority Order"] = np.arange(1, len(wide) + 1)
    if len(wide) != 53:
        raise ValueError("Expected 53 unique shelters after ranking")
    return wide


def prepare_main_table(ranking: pd.DataFrame) -> pd.DataFrame:
    robust = ranking.head(15).copy()
    table = pd.DataFrame()
    table["Robust Priority Order"] = robust["Robust Priority Order"].astype(int)
    table["Municipality"] = robust["Municipality"].map(MUNICIPALITY_LABELS)
    table["Shelter Name"] = robust["Stable Shelter ID"].map(ENGLISH_SHELTER_NAMES)
    table["District"] = (
        robust["District"].astype("string").map(AREA_LABELS).astype("string").fillna("—")
    )
    table["Ward"] = (
        robust["Ward"].astype("string").map(AREA_LABELS).astype("string").fillna("—")
    )
    table["Evacuees"] = robust["Evacuees"].astype(int)
    table["Estimated Female Evacuees"] = robust["Estimated Female Evacuees"].astype(float)
    table["Functional-Support Demand (Low / Base / High)"] = robust.apply(
        lambda row: (
            f"{row['FS Demand Low']:.1f} / {row['FS Demand Base']:.1f} / "
            f"{row['FS Demand High']:.1f}"
        ),
        axis=1,
    )
    table["Estimated Female Functional-Support Evacuees (Base)"] = robust[
        "Estimated Female Functional-Support Evacuees (Base)"
    ].astype(float)
    table["Water Status"] = robust["Water Status"].map(WATER_LABELS).fillna("Not reported")
    table["Temporary Toilets Installed"] = robust["Temporary Toilets Installed"].apply(
        lambda value: "NR" if pd.isna(value) else int(value)
    )
    table["Toilet Cars"] = robust["Toilet Cars"].apply(
        lambda value: "NR" if pd.isna(value) else int(value)
    )
    table["Cross-Scenario Ranks (Low / Base / High)"] = robust.apply(
        lambda row: (
            f"{int(row['Rank Low'])} / {int(row['Rank Base'])} / "
            f"{int(row['Rank High'])}"
        ),
        axis=1,
    )
    table["Worst-Scenario Rank"] = robust["Worst-Scenario Rank"].astype(int)
    if len(table) != 15 or table[list(HEADERS)].isna().any().any():
        raise ValueError("Robust-priority table is incomplete")
    return table[list(HEADERS)]


def write_workbook(table: pd.DataFrame, ranking: pd.DataFrame) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Robust Priorities"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "F2"
    sheet.sheet_view.zoomScale = 75
    sheet.append(list(HEADERS))

    input_order = ranking.sort_values("Stable Shelter ID").reset_index(drop=True)
    for _, row in table.iterrows():
        sheet.append(row.tolist())

    last_row = sheet.max_row
    table_range = f"A1:N{last_row}"
    excel_table = Table(displayName="RobustVerificationPriorities", ref=table_range)
    excel_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(excel_table)

    header_fill = PatternFill("solid", fgColor="27445C")
    header_font = Font(name="Aptos", size=9, bold=True, color="FFFFFF")
    body_font = Font(name="Aptos", size=8.5, color="263746")
    thin_gray = Side(style="thin", color="D8DEE3")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="9EB1C0"))
    sheet.row_dimensions[1].height = 72

    for row in sheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=14):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
        for index in (0, 5, 6, 8, 10, 11, 13):
            row[index].alignment = Alignment(horizontal="right", vertical="center")
        for index in (7, 9, 12):
            row[index].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 39

    for column in ("A", "F", "K", "L", "N"):
        for cell in sheet[column][1:]:
            if isinstance(cell.value, (int, float)) or (
                isinstance(cell.value, str) and cell.value.startswith("=")
            ):
                cell.number_format = "#,##0"
    for column in ("G", "I"):
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0.0"

    water_fills = {
        "Available": "DCEDE5",
        "Partial": "FFF0CC",
        "Unavailable": "F5D7D7",
        "Not reported": "E3E7EA",
    }
    for label, color in water_fills.items():
        sheet.conditional_formatting.add(
            f"J2:J{last_row}",
            FormulaRule(
                formula=[f'LEFT($J2,{len(label)})="{label}"'],
                fill=PatternFill("solid", fgColor=color),
            ),
        )
    for column in ("K", "L"):
        sheet.conditional_formatting.add(
            f"{column}2:{column}{last_row}",
            FormulaRule(
                formula=[f'${column}2="NR"'],
                fill=PatternFill("solid", fgColor="E3E7EA"),
            ),
        )
    priority_fills = ((5, "E7C4BF"), (10, "F2DEC4"), (15, "F7EDCF"))
    lower_bound = 0
    for upper_bound, color in priority_fills:
        sheet.conditional_formatting.add(
            f"A2:A{last_row}",
            FormulaRule(
                formula=[f'AND($A2>{lower_bound},$A2<={upper_bound})'],
                fill=PatternFill("solid", fgColor=color),
            ),
        )
        lower_bound = upper_bound

    widths = {
        "A": 12,
        "B": 16,
        "C": 25,
        "D": 11,
        "E": 11,
        "F": 10,
        "G": 16,
        "H": 23,
        "I": 19,
        "J": 15,
        "K": 13,
        "L": 10,
        "M": 18,
        "N": 13,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet["H1"].comment = Comment(
        "Scenario order is Low / Base / High. Values are estimated evacuees with "
        "functional-support needs under the declared overlap assumptions.",
        "Mike Li",
    )
    sheet["M1"].comment = Comment(
        "Cross-city ranks use descending estimated functional-support demand, "
        "female functional-support demand, and observed evacuees, followed by the "
        "stable shelter identifier.",
        "Mike Li",
    )
    sheet["N1"].comment = Comment(
        "Worst-scenario rank M_j is the maximum of the three cross-city scenario ranks.",
        "Mike Li",
    )
    sheet.auto_filter.ref = table_range
    sheet.print_area = table_range
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.18
    sheet.page_margins.right = 0.18
    sheet.page_margins.top = 0.3
    sheet.page_margins.bottom = 0.3

    inputs = workbook.create_sheet("Ranking Inputs")
    inputs.sheet_view.showGridLines = False
    input_headers = (
        "Stable Shelter ID",
        "Municipality",
        "Shelter Number",
        "Shelter Name",
        "Evacuees",
        "Estimated Female Evacuees",
        "FS Demand Low",
        "FS Demand Base",
        "FS Demand High",
        "Female FS Demand Base",
        "Rank Low",
        "Rank Base",
        "Rank High",
        "Worst-Scenario Rank",
        "Rank Span",
        "Location Resolution",
    )
    inputs.append(list(input_headers))
    for row_number, (_, row) in enumerate(input_order.iterrows(), start=2):
        inputs.append(
            [
                str(row["Stable Shelter ID"]),
                MUNICIPALITY_LABELS[str(row["Municipality"])],
                int(row["Shelter Number"]),
                ENGLISH_SHELTER_NAMES[str(row["Stable Shelter ID"])],
                int(row["Evacuees"]),
                float(row["Estimated Female Evacuees"]),
                float(row["FS Demand Low"]),
                float(row["FS Demand Base"]),
                float(row["FS Demand High"]),
                float(row["Estimated Female Functional-Support Evacuees (Base)"]),
                int(row["Rank Low"]),
                int(row["Rank Base"]),
                int(row["Rank High"]),
                f"=MAX(K{row_number}:M{row_number})",
                f"=MAX(K{row_number}:M{row_number})-MIN(K{row_number}:M{row_number})",
                str(row["Location Resolution"]),
            ]
        )
    for cell in inputs[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    inputs.row_dimensions[1].height = 58
    for row in inputs.iter_rows(min_row=2, max_row=inputs.max_row, min_col=1, max_col=16):
        for cell in row:
            cell.font = Font(name="Aptos", size=8, color="263746")
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.border = Border(bottom=thin_gray)
        for index in (0, 1, 3, 15):
            row[index].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for column in ("F", "G", "H", "I", "J"):
        for cell in inputs[column][1:]:
            cell.number_format = "#,##0.0"
    input_widths = {
        "A": 12,
        "B": 16,
        "C": 10,
        "D": 26,
        "E": 10,
        "F": 15,
        "G": 13,
        "H": 13,
        "I": 13,
        "J": 16,
        "K": 10,
        "L": 10,
        "M": 10,
        "N": 13,
        "O": 10,
        "P": 24,
    }
    for column, width in input_widths.items():
        inputs.column_dimensions[column].width = width
    inputs.freeze_panes = "E2"
    inputs.sheet_view.zoomScale = 75
    inputs.print_area = f"A1:P{inputs.max_row}"
    inputs.print_title_rows = "1:1"
    inputs.page_setup.orientation = "landscape"
    inputs.page_setup.paperSize = inputs.PAPERSIZE_A3
    inputs.page_setup.fitToWidth = 1
    inputs.page_setup.fitToHeight = 0
    inputs.sheet_properties.pageSetUpPr.fitToPage = True

    notes = workbook.create_sheet("Notes")
    notes.sheet_view.showGridLines = False
    notes["A1"] = "Ranking definitions, evidence boundaries, and sources"
    notes["A1"].font = Font(name="Aptos Display", size=15, bold=True, color="27445C")
    notes.append([])
    notes.append(["Item", "Definition or boundary", "Source URL"])
    note_rows = (
        (
            "Robust-set rule",
            "Rank all 53 matched shelters separately under Low, Base, and High. Retain the 15 shelters with the smallest worst-scenario rank; break ties by Base functional-support demand, Base female functional-support demand, evacuees, then stable shelter ID.",
            "",
        ),
        (
            "Interpretation boundary",
            "Ranks prioritize field verification only. They do not estimate intervention benefit, prove a toilet deficit, or optimize deployment.",
            "",
        ),
        (
            "Synthetic subgroup demand",
            "Female and functional-support values apply residential catchment composition to observed shelter occupancy; they are not observed shelter-user composition.",
            "https://www.e-stat.go.jp/gis/statmap-search?page=1&type=1&toukeiCode=00200521",
        ),
        (
            "Kumamoto shelter evidence",
            "The matched shelter roster is the 12 August 2026 12:00 city snapshot. Water and deployment fields are incomplete and NR is not zero.",
            "https://www.city.kumamoto.jp/kiji00372080/3_72080_514016_up_mh1811lq.pdf",
        ),
        (
            "Yatsushiro shelter evidence",
            "The matched shelter table is the 12 August 2026 12:00 snapshot. Reported temporary toilets are not equivalent to verified total functional capacity.",
            "https://www.city.yatsushiro.lg.jp/kiji00326750/3_26750_158214_up_bjxkz7dq.pdf",
        ),
        (
            "Toilet cars",
            "Vehicle counts remain vehicles and are not converted to stalls without vehicle-specific evidence.",
            "",
        ),
    )
    for values in note_rows:
        notes.append(values)
    for cell in notes[3]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
    notes.row_dimensions[3].height = 24
    for row in notes.iter_rows(min_row=4, max_row=notes.max_row, min_col=1, max_col=3):
        row[0].font = Font(name="Aptos", size=9, bold=True, color="263746")
        row[1].font = Font(name="Aptos", size=8.5, color="263746")
        row[2].font = Font(name="Aptos", size=8, color="0563C1", underline="single")
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
        notes.row_dimensions[row[0].row].height = 42
    notes.column_dimensions["A"].width = 27
    notes.column_dimensions["B"].width = 96
    notes.column_dimensions["C"].width = 78
    notes.freeze_panes = "A4"
    notes.sheet_view.zoomScale = 85
    notes.print_area = f"A1:C{notes.max_row}"
    notes.page_setup.orientation = "landscape"
    notes.page_setup.paperSize = notes.PAPERSIZE_A4
    notes.page_setup.fitToWidth = 1
    notes.page_setup.fitToHeight = 1
    notes.sheet_properties.pageSetUpPr.fitToPage = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)


def verify_workbook(expected: pd.DataFrame) -> None:
    workbook = load_workbook(OUTPUT, data_only=False, read_only=False)
    if workbook.sheetnames != ["Robust Priorities", "Ranking Inputs", "Notes"]:
        raise ValueError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    sheet = workbook["Robust Priorities"]
    if sheet.max_row != 16 or sheet.max_column != 14:
        raise ValueError(f"Unexpected main-table dimensions: {sheet.max_row} x {sheet.max_column}")
    if tuple(cell.value for cell in sheet[1]) != HEADERS:
        raise ValueError("Main-table headers do not match the planned 14-column table")
    if any(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for row in sheet.iter_rows()
        for cell in row
    ):
        raise ValueError("The article-facing Robust Priorities sheet must contain static values")
    if [sheet.cell(row=row, column=14).value for row in range(2, 17)] != expected[
        "Worst-Scenario Rank"
    ].tolist():
        raise ValueError("Worst-scenario ranks changed in the main table")
    inputs = workbook["Ranking Inputs"]
    if inputs.max_row != 54 or inputs.max_column != 16:
        raise ValueError("Ranking Inputs must contain 53 shelters and 16 columns")
    if not all(str(inputs.cell(row=row, column=14).value).startswith("=MAX") for row in range(2, 55)):
        raise ValueError("Worst-rank formulas are missing from Ranking Inputs")
    if tuple(expected["Robust Priority Order"].astype(int)) != tuple(range(1, 16)):
        raise ValueError("Robust priority order is not sequential")
    japanese_cells = []
    for workbook_sheet in workbook.worksheets:
        for row in workbook_sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and JAPANESE_PATTERN.search(cell.value):
                    japanese_cells.append(f"{workbook_sheet.title}!{cell.coordinate}")
    if japanese_cells:
        raise ValueError(f"Japanese text remains in workbook: {japanese_cells[:5]}")
    workbook.close()


def main() -> None:
    source = pd.read_parquet(SOURCE)
    ranking = build_ranking(source)
    table = prepare_main_table(ranking)
    write_workbook(table, ranking)
    verify_workbook(table)
    selected = ranking.head(15)
    city_counts = selected["Municipality"].map(MUNICIPALITY_LABELS).value_counts().to_dict()
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(f"Robust set: 15 shelters from 53; city distribution: {city_counts}")
    print(
        "Worst-scenario rank range: "
        f"{int(selected['Worst-Scenario Rank'].min())}-"
        f"{int(selected['Worst-Scenario Rank'].max())}"
    )


if __name__ == "__main__":
    main()
