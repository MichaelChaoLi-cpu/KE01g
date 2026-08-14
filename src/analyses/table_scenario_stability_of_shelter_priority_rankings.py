#!/usr/bin/env python3
"""Scenario Stability of Shelter Priority Rankings.

Plan: Report within-city shelter verification ranks under the low, base, and
high overlap scenarios, together with rank movement and robust high-priority
membership for the complete matched shelter sample.
Framework: Section 5 scenario-stability contrast; Section 6 within-city
lexicographic rank pi, top-20-percent membership H, and rank span B; Section 7
verification-priority robustness workflow. Ranks guide verification only and
do not estimate intervention benefit or optimize deployment.
"""

from __future__ import annotations

from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from table_shelter_level_equity_demand_scenario_estimates import (
    AREA_LABELS,
    ENGLISH_SHELTER_NAMES,
    LOCATION_LABELS,
    MUNICIPALITY_LABELS,
    MUNICIPALITY_PREFIXES,
)


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/processed/shelter_equity_scenarios_preprocessed.parquet"
OUTPUT = (
    ROOT
    / "data/results/tables/"
    "Table_scenario_stability_of_shelter_priority_rankings.xlsx"
)

SCENARIOS = ("low", "base", "high")
JAPANESE_PATTERN = re.compile(r"[\u3040-\u30ff\u3400-\u9fff]")
HEADERS = (
    "Stable Shelter ID",
    "Municipality",
    "Shelter Name",
    "Area",
    "Evacuees",
    "Low Rank",
    "Base Rank",
    "High Rank",
    "Maximum Rank Change",
    "Priority Stability Classification",
)


def stable_id(frame: pd.DataFrame) -> pd.Series:
    prefixes = frame["Municipality"].map(MUNICIPALITY_PREFIXES)
    if prefixes.isna().any():
        raise ValueError("Unexpected municipality in matched shelter sample")
    return prefixes + frame["Shelter Number"].astype(int).astype(str).str.zfill(2)


def rank_shelters(frame: pd.DataFrame) -> pd.DataFrame:
    required = {
        "Municipality",
        "Shelter Number",
        "Shelter Name",
        "District",
        "Ward",
        "Scenario",
        "Evacuees",
        "Estimated Functional Support Evacuees",
        "Estimated Female Functional Support Evacuees",
        "Location Resolution",
    }
    missing = required.difference(frame.columns)
    if missing:
        raise KeyError(f"Missing required columns: {sorted(missing)}")
    frame = frame.loc[
        frame["Municipality"].isin(MUNICIPALITY_LABELS)
        & frame["Scenario"].astype("string").isin(SCENARIOS)
    ].copy()
    if len(frame) != 159:
        raise ValueError("Expected 53 shelters under three scenarios")
    if frame.duplicated(["Municipality", "Shelter Number", "Scenario"]).any():
        raise ValueError("Shelter-scenario keys must be unique")
    if frame[list(required)].drop(columns=["District", "Ward"]).isna().any().any():
        raise ValueError("Ranking inputs must be complete")

    frame["Stable Shelter ID"] = stable_id(frame)
    ranked_parts: list[pd.DataFrame] = []
    for municipality, city in frame.groupby("Municipality", sort=False):
        for scenario in SCENARIOS:
            part = city.loc[city["Scenario"].astype(str).eq(scenario)].copy()
            part = part.sort_values(
                [
                    "Estimated Functional Support Evacuees",
                    "Estimated Female Functional Support Evacuees",
                    "Evacuees",
                    "Stable Shelter ID",
                ],
                ascending=[False, False, False, True],
                kind="stable",
            )
            part["Verification Rank"] = np.arange(1, len(part) + 1)
            ranked_parts.append(part)
    ranked = pd.concat(ranked_parts, ignore_index=True)

    base_columns = [
        "Stable Shelter ID",
        "Municipality",
        "Shelter Number",
        "Shelter Name",
        "District",
        "Ward",
        "Evacuees",
        "Location Resolution",
    ]
    base = ranked.loc[ranked["Scenario"].astype(str).eq("base"), base_columns].copy()
    rank_wide = (
        ranked.pivot(
            index="Stable Shelter ID", columns="Scenario", values="Verification Rank"
        )
        .rename(columns={s: f"{s.title()} Rank" for s in SCENARIOS})
        .reset_index()
    )
    wide = base.merge(rank_wide, on="Stable Shelter ID", validate="one_to_one")
    rank_columns = ["Low Rank", "Base Rank", "High Rank"]
    wide["Maximum Rank Change"] = (
        wide[rank_columns].max(axis=1) - wide[rank_columns].min(axis=1)
    ).astype(int)
    wide["Top Threshold"] = (
        np.ceil(0.20 * wide.groupby("Municipality")["Stable Shelter ID"].transform("size"))
        .astype(int)
    )
    wide["Stable High Priority"] = wide[rank_columns].max(axis=1).le(
        wide["Top Threshold"]
    )
    wide["Sensitivity"] = pd.cut(
        wide["Maximum Rank Change"],
        bins=[-np.inf, 0, 2, np.inf],
        labels=["Stable", "Low sensitivity", "Scenario-sensitive"],
    ).astype(str)
    wide["Priority Stability Classification"] = np.where(
        wide["Stable High Priority"],
        "Stable high priority",
        wide["Sensitivity"],
    )

    expected_counts = {"八代市": 38, "熊本市": 15}
    if wide.groupby("Municipality").size().to_dict() != expected_counts:
        raise ValueError("Unexpected municipality shelter counts")
    for municipality, count in expected_counts.items():
        city = wide.loc[wide["Municipality"].eq(municipality)]
        for column in rank_columns:
            if sorted(city[column].astype(int).tolist()) != list(range(1, count + 1)):
                raise ValueError(f"Invalid rank permutation: {municipality}, {column}")
    return wide.sort_values(
        ["Municipality", "Base Rank", "Stable Shelter ID"],
        key=lambda values: (
            values.map({"八代市": 0, "熊本市": 1})
            if values.name == "Municipality"
            else values
        ),
        kind="stable",
    ).reset_index(drop=True)


def prepare_table(ranking: pd.DataFrame) -> pd.DataFrame:
    table = pd.DataFrame()
    table["Stable Shelter ID"] = ranking["Stable Shelter ID"]
    table["Municipality"] = ranking["Municipality"].map(MUNICIPALITY_LABELS)
    table["Shelter Name"] = ranking["Stable Shelter ID"].map(ENGLISH_SHELTER_NAMES)
    source_area = ranking["District"].combine_first(ranking["Ward"])
    table["Area"] = source_area.map(AREA_LABELS)
    table["Evacuees"] = ranking["Evacuees"].astype(int)
    for scenario in ("Low", "Base", "High"):
        table[f"{scenario} Rank"] = ranking[f"{scenario} Rank"].astype(int)
    table["Maximum Rank Change"] = ranking["Maximum Rank Change"].astype(int)
    table["Priority Stability Classification"] = ranking[
        "Priority Stability Classification"
    ]
    if table[list(HEADERS)].isna().any().any():
        raise ValueError("English table contains missing values")
    return table[list(HEADERS)]


def write_workbook(table: pd.DataFrame, ranking: pd.DataFrame) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ranking Stability"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "E2"
    sheet.sheet_view.zoomScale = 85
    sheet.append(list(HEADERS))

    for _, row in table.iterrows():
        sheet.append(row.tolist())

    last_row = sheet.max_row
    excel_table = Table(displayName="ShelterRankingStability", ref=f"A1:J{last_row}")
    excel_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(excel_table)

    navy = "27445C"
    text = "263746"
    border = Side(style="thin", color="D8DEE3")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="9EB1C0"))
    sheet.row_dimensions[1].height = 52

    for row in sheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=10):
        for cell in row:
            cell.font = Font(name="Aptos", size=8.5, color=text)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=border)
        for index in (4, 5, 6, 7):
            row[index].alignment = Alignment(horizontal="right", vertical="center")
        row[8].alignment = Alignment(horizontal="center", vertical="center")
        row[9].alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True, indent=1
        )
        row[9].border = Border(
            left=Side(style="thin", color="BCC8D1"), bottom=border
        )
        sheet.row_dimensions[row[0].row].height = 31

    sheet.conditional_formatting.add(
        f"J2:J{last_row}",
        FormulaRule(
            formula=['$J2="Stable high priority"'],
            fill=PatternFill("solid", fgColor="F3CEC5"),
            font=Font(name="Aptos", size=8.5, bold=True, color="8E342A"),
        ),
    )
    sheet.conditional_formatting.add(
        f"J2:J{last_row}",
        FormulaRule(
            formula=['$J2="Scenario-sensitive"'],
            fill=PatternFill("solid", fgColor="DDD1E8"),
            font=Font(name="Aptos", size=8.5, bold=True, color="64477B"),
        ),
    )
    sheet.conditional_formatting.add(
        f"I2:I{last_row}",
        CellIsRule(
            operator="greaterThanOrEqual",
            formula=["3"],
            fill=PatternFill("solid", fgColor="DDD1E8"),
        ),
    )

    widths = {
        "A": 13,
        "B": 16,
        "C": 36,
        "D": 16,
        "E": 10,
        "F": 9,
        "G": 9,
        "H": 9,
        "I": 15,
        "J": 28,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = f"A1:J{last_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 2
    sheet.print_title_rows = "1:1"
    sheet.print_area = f"A1:J{last_row}"
    sheet.sheet_properties.outlinePr.summaryBelow = True
    # Visually separate the two cities while retaining one sortable table.
    kumamoto_start = int(table.index[table["Municipality"].eq("Kumamoto City")][0]) + 2
    for cell in sheet[kumamoto_start]:
        cell.border = Border(top=Side(style="medium", color="6F8799"), bottom=border)

    notes = workbook.create_sheet("Notes")
    notes.sheet_view.showGridLines = False
    notes.column_dimensions["A"].width = 31
    notes.column_dimensions["B"].width = 110
    notes.append(["Field", "Definition"])
    note_rows = [
        ("Ranking universe", "Within each city, all 38 Yatsushiro and 15 Kumamoto shelters are ranked separately under every scenario."),
        ("Ranking rule", "Descending Estimated Functional-Support Evacuees, then Estimated Female Functional-Support Evacuees, Evacuees, and Stable Shelter ID."),
        ("Stable high priority", "The shelter remains within the city's top 20% under all three overlap scenarios: top 8 in Yatsushiro or top 3 in Kumamoto."),
        ("Maximum rank change", "Largest minus smallest within-city rank across Low, Base, and High scenarios."),
        ("Stability classes", "Stable = no rank change; Low sensitivity = change of 1–2 places; Scenario-sensitive = change of at least 3 places. Stable high priority overrides these display labels."),
        ("Interpretation", "Synthetic overlap-scenario ranks guide field verification. They do not measure observed subgroup composition, intervention benefit, or optimal deployment."),
        ("Area", "Yatsushiro district or Kumamoto ward, combined because only one field applies to each shelter."),
        ("Location evidence", f"{int(ranking['Location Resolution'].eq('exact shelter master match').sum())} exact shelter matches and {int(ranking['Location Resolution'].eq('district anchor fallback').sum())} district-anchor fallbacks in the complete sample."),
    ]
    for row in note_rows:
        notes.append(row)
    for cell in notes[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="left", vertical="center")
    for row in notes.iter_rows(min_row=2, max_row=notes.max_row, min_col=1, max_col=2):
        row[0].font = Font(name="Aptos", size=9, bold=True, color=text)
        row[1].font = Font(name="Aptos", size=9, color=text)
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = Border(bottom=border)
        notes.row_dimensions[row[0].row].height = 42
    notes.freeze_panes = "A2"
    notes.sheet_properties.pageSetUpPr.fitToPage = True
    notes.page_setup.orientation = "landscape"
    notes.page_setup.paperSize = notes.PAPERSIZE_A4
    notes.page_setup.fitToWidth = 1
    notes.page_setup.fitToHeight = 1
    notes.print_area = f"A1:B{notes.max_row}"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)


def validate_output(table: pd.DataFrame) -> None:
    workbook = load_workbook(OUTPUT, data_only=False)
    if workbook.sheetnames != ["Ranking Stability", "Notes"]:
        raise ValueError("Unexpected workbook sheet structure")
    sheet = workbook["Ranking Stability"]
    if sheet.max_row != 54 or sheet.max_column != 10:
        raise ValueError("Unexpected main-table dimensions")
    formula_count = sum(
        1
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )
    if formula_count != 0:
        raise ValueError("The article-facing Ranking Stability sheet must contain static values")
    japanese_cells = []
    for workbook_sheet in workbook.worksheets:
        for row in workbook_sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and JAPANESE_PATTERN.search(cell.value):
                    japanese_cells.append(f"{workbook_sheet.title}!{cell.coordinate}")
    if japanese_cells:
        raise ValueError(f"Japanese text remains in workbook: {japanese_cells[:5]}")
    if len(table) != 53:
        raise ValueError("Expected 53 shelter rows")


def main() -> None:
    ranking = rank_shelters(pd.read_parquet(SOURCE))
    table = prepare_table(ranking)
    write_workbook(table, ranking)
    validate_output(table)
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print("Main table: 53 shelters x 10 columns")
    print(
        "Stable high priority: "
        f"{int(ranking['Stable High Priority'].sum())}; "
        "scenario-sensitive: "
        f"{int(ranking['Sensitivity'].eq('Scenario-sensitive').sum())}"
    )


if __name__ == "__main__":
    main()
