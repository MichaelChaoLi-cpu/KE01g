#!/usr/bin/env python3
"""Scenario Calibration, Catchment, and Evidence Audit.

Plan: Audit residential population assignment, physical-disability and
long-term-care calibration, certification sex reconciliation, near-time city
calibration, shelter location linkage, and interpretation boundaries.
Framework: Section 5 two-tier evidence design; Section 6 bounded mesh
allocation and linkage coverage Q; Section 7 prerequisite calibration and
evidence audit. Passing an audit permits scenario analysis only at the stated
resolution and never establishes observed subgroup composition or adequacy.
"""

from __future__ import annotations

import json
from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from table_municipality_data_coverage_and_evidence_resolution_audit import (
    LEVEL_COLUMNS,
    read_and_validate_inputs,
)


ROOT = Path(__file__).resolve().parents[2]
QA_SOURCE = ROOT / "data/exp/scenario-planning/qa_summary.json"
MUNICIPALITY_SOURCE = (
    ROOT
    / "data/processed/"
    "kumamoto_prefecture_municipality_shelter_totals_2026-08-12_1400_preprocessed.parquet"
)
MESH_SOURCE = ROOT / "data/processed/mesh_equity_scenarios_preprocessed.parquet"
SHELTER_SOURCE = ROOT / "data/processed/shelter_equity_scenarios_preprocessed.parquet"
OUTPUT = (
    ROOT
    / "data/results/tables/"
    "Table_scenario_calibration_catchment_and_evidence_audit.xlsx"
)

JAPANESE_PATTERN = re.compile(r"[\u3040-\u30ff\u3400-\u9fff]")
HEADERS = (
    "Audit Component",
    "Target Value",
    "Allocated or Observed Value",
    "Difference",
    "Residential Population Boundary Difference",
    "Male + Female Reconciliation",
    "Exact Location Count",
    "District-Anchor Fallback Count",
    "Location Resolution",
    "Reference Period",
    "Permitted Use",
    "Prohibited Interpretation",
)

CITY_LABELS = {"熊本市": "Kumamoto City", "八代市": "Yatsushiro City"}
PUBLISHED_POPULATION = {"熊本市": 738865.0, "八代市": 123067.0}
PHYSICAL_DISABILITY_TARGET = {"熊本市": 26726.0, "八代市": 5296.0}
LONG_TERM_CARE_TARGET = {"熊本市": 41944.0, "八代市": 8698.0}


def clean_number(value: float, tolerance: float = 1e-7) -> int | float:
    if abs(value - round(value)) <= tolerance:
        return int(round(value))
    return float(value)


def build_audit() -> pd.DataFrame:
    with QA_SOURCE.open(encoding="utf-8") as handle:
        qa = json.load(handle)
    municipality = pd.read_parquet(MUNICIPALITY_SOURCE)
    mesh = pd.read_parquet(MESH_SOURCE)
    shelter = pd.read_parquet(SHELTER_SOURCE)
    _, certification = read_and_validate_inputs()

    if municipality["Municipality Observation Time"].nunique() != 1:
        raise ValueError("Municipality observation time is not common")
    if int(municipality["Open Shelter Count"].sum()) != 81:
        raise ValueError("Municipality records do not reproduce 81 open shelters")
    if int(municipality["Evacuees"].sum()) != 3585:
        raise ValueError("Municipality records do not reproduce 3,585 evacuees")
    base_mesh = mesh.loc[mesh["Scenario"].astype(str).eq("base")].copy()
    base_shelter = shelter.loc[shelter["Scenario"].astype(str).eq("base")].copy()
    if len(base_shelter) != 53:
        raise ValueError("Expected 53 shelters in the matched base scenario")

    reconciliation_checks = 0
    for level in LEVEL_COLUMNS:
        matches = (
            certification[f"Total {level}"]
            == certification[f"Male {level}"] + certification[f"Female {level}"]
        )
        reconciliation_checks += int(matches.sum())
    expected_reconciliation_checks = len(certification) * len(LEVEL_COLUMNS)
    if reconciliation_checks != expected_reconciliation_checks:
        raise ValueError("Certification sex reconciliation failed")

    rows: list[dict[str, object]] = []
    for city in ("熊本市", "八代市"):
        allocated = float(
            base_mesh.loc[base_mesh["Municipality"].eq(city), "Total Population"].sum()
        )
        target = PUBLISHED_POPULATION[city]
        difference = allocated - target
        qa_difference = float(qa["mesh_boundary_population_difference"][city])
        if not np.isclose(difference, qa_difference):
            raise ValueError(f"Residential boundary audit mismatch for {city}")
        rows.append(
            {
                "Audit Component": f"{CITY_LABELS[city]} residential assignment",
                "Target Value": clean_number(target),
                "Allocated or Observed Value": clean_number(allocated),
                "Difference": clean_number(difference),
                "Residential Population Boundary Difference": clean_number(difference),
                "Male + Female Reconciliation": "N/A",
                "Exact Location Count": "N/A",
                "District-Anchor Fallback Count": "N/A",
                "Location Resolution": "Mesh centers in 2025 boundary",
                "Reference Period": "2020 census; 2025 boundary",
                "Permitted Use": "Residential scenario input",
                "Prohibited Interpretation": "Observed evacuee composition",
            }
        )

    calibration_specs = (
        (
            "Physical-disability",
            "Estimated Physical Disability Population",
            PHYSICAL_DISABILITY_TARGET,
            {"熊本市": "FY2024", "八代市": "FY2024"},
            "Observed disability prevalence",
        ),
        (
            "Long-term-care",
            "Estimated Long-Term Care Population",
            LONG_TERM_CARE_TARGET,
            {"熊本市": "2025-03", "八代市": "FY2026 projection"},
            "Observed LTC prevalence",
        ),
    )
    for label, column, target_map, reference_map, prohibited in calibration_specs:
        for city in ("熊本市", "八代市"):
            target = target_map[city]
            allocated = float(
                base_mesh.loc[base_mesh["Municipality"].eq(city), column].sum()
            )
            difference = allocated - target
            if not np.isclose(difference, 0, atol=1e-7):
                raise ValueError(f"{label} calibration failed for {city}: {difference}")
            rows.append(
                {
                    "Audit Component": f"{CITY_LABELS[city]} {label.lower()} calibration",
                    "Target Value": clean_number(target),
                    "Allocated or Observed Value": clean_number(allocated),
                    "Difference": clean_number(difference),
                    "Residential Population Boundary Difference": "N/A",
                    "Male + Female Reconciliation": "N/A",
                    "Exact Location Count": "N/A",
                    "District-Anchor Fallback Count": "N/A",
                    "Location Resolution": "Bounded mesh allocation",
                    "Reference Period": reference_map[city],
                    "Permitted Use": "Synthetic need allocation",
                    "Prohibited Interpretation": prohibited,
                }
            )

    rows.append(
        {
            "Audit Component": "Long-term-care sex reconciliation",
            "Target Value": expected_reconciliation_checks,
            "Allocated or Observed Value": reconciliation_checks,
            "Difference": reconciliation_checks - expected_reconciliation_checks,
            "Residential Population Boundary Difference": "N/A",
            "Male + Female Reconciliation": "Exact: 88 of 88 checks",
            "Exact Location Count": "N/A",
            "District-Anchor Fallback Count": "N/A",
            "Location Resolution": "11 municipalities; 8 fields",
            "Reference Period": "April 2026",
            "Permitted Use": "Sex-specific LTC scenarios",
            "Prohibited Interpretation": "Current shelter composition",
        }
    )

    municipality_by_city = municipality.set_index("Municipality")
    for city in ("熊本市", "八代市"):
        target = int(municipality_by_city.loc[CITY_LABELS[city], "Evacuees"])
        city_shelters = base_shelter.loc[base_shelter["Municipality"].eq(city)]
        observed = int(city_shelters["Evacuees"].sum())
        exact = int(city_shelters["Location Resolution"].astype(str).eq("exact shelter master match").sum())
        fallback = int(city_shelters["Location Resolution"].astype(str).eq("district anchor fallback").sum())
        target_shelters = int(municipality_by_city.loc[CITY_LABELS[city], "Open Shelter Count"])
        observed_shelters = len(city_shelters)
        rows.append(
            {
                "Audit Component": (
                    f"{CITY_LABELS[city]} near-time shelter calibration "
                    f"({target_shelters} vs {observed_shelters} shelters)"
                ),
                "Target Value": target,
                "Allocated or Observed Value": observed,
                "Difference": observed - target,
                "Residential Population Boundary Difference": "N/A",
                "Male + Female Reconciliation": "N/A",
                "Exact Location Count": exact,
                "District-Anchor Fallback Count": fallback,
                "Location Resolution": "Exact" if fallback == 0 else "Mixed exact + district anchor",
                "Reference Period": "Municipality 14:00; shelters 12:00",
                "Permitted Use": "Near-time shelter screening",
                "Prohibited Interpretation": "Exact 14:00 shelter distribution",
            }
        )

    for city in ("熊本市", "八代市"):
        city_shelters = base_shelter.loc[base_shelter["Municipality"].eq(city)]
        target = len(city_shelters)
        exact = int(city_shelters["Location Resolution"].astype(str).eq("exact shelter master match").sum())
        fallback = int(city_shelters["Location Resolution"].astype(str).eq("district anchor fallback").sum())
        located = exact + fallback
        rows.append(
            {
                "Audit Component": f"{CITY_LABELS[city]} shelter-location linkage",
                "Target Value": target,
                "Allocated or Observed Value": located,
                "Difference": located - target,
                "Residential Population Boundary Difference": "N/A",
                "Male + Female Reconciliation": "N/A",
                "Exact Location Count": exact,
                "District-Anchor Fallback Count": fallback,
                "Location Resolution": "Exact" if fallback == 0 else "Mixed exact + district anchor",
                "Reference Period": "2026 shelter roster; 2025 master",
                "Permitted Use": "Exact-site mapping" if fallback == 0 else "District-level screening",
                "Prohibited Interpretation": "Unverified facility coordinates" if fallback == 0 else "Exact fallback-site placement",
            }
        )

    audit = pd.DataFrame(rows, columns=HEADERS)
    if len(audit) != 11 or audit[list(HEADERS)].isna().any().any():
        raise ValueError("Audit must contain 11 complete components")
    expected_location_counts = qa["location_resolution_counts"]
    if exact_location_total(audit) != int(expected_location_counts["exact shelter master match"]):
        raise ValueError("Exact-location count does not reproduce QA summary")
    if fallback_location_total(audit) != int(expected_location_counts["district anchor fallback"]):
        raise ValueError("Fallback-location count does not reproduce QA summary")
    return audit


def exact_location_total(audit: pd.DataFrame) -> int:
    linkage = audit["Audit Component"].str.endswith("shelter-location linkage")
    return int(pd.to_numeric(audit.loc[linkage, "Exact Location Count"]).sum())


def fallback_location_total(audit: pd.DataFrame) -> int:
    linkage = audit["Audit Component"].str.endswith("shelter-location linkage")
    return int(pd.to_numeric(audit.loc[linkage, "District-Anchor Fallback Count"]).sum())


def write_workbook(audit: pd.DataFrame) -> None:
    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    sheet = workbook.active
    sheet.title = "Calibration Audit"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "B2"
    sheet.sheet_view.zoomScale = 72
    sheet.append(list(HEADERS))
    for row_number, (_, row) in enumerate(audit.iterrows(), start=2):
        values: list[object] = [row["Audit Component"]]
        for column_index in range(2, 13):
            input_column = chr(ord("A") + column_index - 1)
            values.append(f"='Audit Inputs'!{input_column}{row_number}")
        sheet.append(values)

    last_row = sheet.max_row
    excel_table = Table(displayName="ScenarioCalibrationAudit", ref=f"A1:L{last_row}")
    excel_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(excel_table)
    navy = "27445C"
    text = "263746"
    border = Side(style="thin", color="D8DEE3")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=8.5, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="9EB1C0"))
    sheet.row_dimensions[1].height = 70
    for row in sheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=12):
        for column_index, cell in enumerate(row):
            cell.font = Font(name="Aptos", size=8, color=text)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(
                left=border if column_index > 0 else Side(style=None), bottom=border
            )
        for index in (1, 2, 3, 4, 6, 7):
            row[index].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 46
    for audit_index, (_, audit_row) in enumerate(audit.iterrows(), start=2):
        for column_index in (2, 3, 4, 5, 7, 8):
            if audit_row.iloc[column_index - 1] == "N/A":
                sheet.cell(audit_index, column_index).alignment = Alignment(
                    horizontal="center", vertical="center"
                )

    green = PatternFill("solid", fgColor="DCEDE5")
    amber = PatternFill("solid", fgColor="FFF0CC")
    red = PatternFill("solid", fgColor="F5D7D7")
    sheet.conditional_formatting.add(
        f"D2:D{last_row}",
        FormulaRule(formula=["$D2=0"], fill=green),
    )
    sheet.conditional_formatting.add(
        f"D2:D{last_row}",
        FormulaRule(formula=['AND(ISNUMBER($D2),$D2<>0)'], fill=amber),
    )
    sheet.conditional_formatting.add(
        f"F2:F{last_row}",
        FormulaRule(formula=['LEFT($F2,5)="Exact"'], fill=green),
    )
    sheet.conditional_formatting.add(
        f"I2:I{last_row}",
        FormulaRule(formula=['LEFT($I2,5)="Mixed"'], fill=amber),
    )
    sheet.conditional_formatting.add(f"K2:K{last_row}", FormulaRule(formula=["TRUE"], fill=green))
    sheet.conditional_formatting.add(f"L2:L{last_row}", FormulaRule(formula=["TRUE"], fill=red))
    for column in ("B", "C", "D", "E", "G", "H"):
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0;-#,##0;0"
    widths = {
        "A": 31,
        "B": 13,
        "C": 17,
        "D": 12,
        "E": 19,
        "F": 20,
        "G": 12,
        "H": 17,
        "I": 26,
        "J": 24,
        "K": 25,
        "L": 29,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = f"A1:L{last_row}"

    inputs = workbook.create_sheet("Audit Inputs")
    inputs.sheet_view.showGridLines = False
    inputs.append(list(HEADERS))
    for _, row in audit.iterrows():
        inputs.append(row.tolist())
    for cell in inputs[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=8.5, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    inputs.row_dimensions[1].height = 62
    for row in inputs.iter_rows(min_row=2, max_row=inputs.max_row):
        for cell in row:
            cell.font = Font(name="Aptos", size=8, color=text)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=border)
        inputs.row_dimensions[row[0].row].height = 42
    for column, width in widths.items():
        inputs.column_dimensions[column].width = width
    inputs.freeze_panes = "B2"

    notes = workbook.create_sheet("Notes")
    notes.sheet_view.showGridLines = False
    notes.column_dimensions["A"].width = 31
    notes.column_dimensions["B"].width = 108
    notes.append(["Field", "Definition"])
    note_rows = [
        ("Difference", "Allocated or observed value minus target. Small residential differences arise because mesh centers are assigned to 2025 boundaries; nonzero near-time shelter differences arise because the shelter and municipality snapshots are two hours apart."),
        ("Calibration", "Bounded mesh allocation must reproduce each city administrative physical-disability and long-term-care input within numerical tolerance."),
        ("Sex reconciliation", "All-sex long-term-care totals must equal male plus female counts for 11 municipalities across Support Levels 1–2, Care Levels 1–5, and Certified Total: 88 checks."),
        ("Location resolution", "Exact matches support facility mapping. District-anchor fallbacks support district screening only and cannot be treated as verified facility coordinates."),
        ("Scenario boundary", "Residential, physical-disability, and long-term-care inputs generate synthetic planning scenarios. They do not observe the sex, disability, care, or accessible-toilet needs of current evacuees."),
        ("Service boundary", "Passing these checks does not establish toilet quantity adequacy, operability, accessibility, privacy, safety, menstrual-hygiene support, or wastewater functionality."),
    ]
    for row in note_rows:
        notes.append(row)
    for cell in notes[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
    for row in notes.iter_rows(min_row=2, max_row=notes.max_row):
        row[0].font = Font(name="Aptos", size=9, bold=True, color=text)
        row[1].font = Font(name="Aptos", size=9, color=text)
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = Border(bottom=border)
        notes.row_dimensions[row[0].row].height = 48
    notes.freeze_panes = "A2"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)


def validate_output(audit: pd.DataFrame) -> None:
    workbook = load_workbook(OUTPUT, data_only=False)
    if workbook.sheetnames != ["Calibration Audit", "Audit Inputs", "Notes"]:
        raise ValueError("Unexpected workbook sheet structure")
    sheet = workbook["Calibration Audit"]
    if sheet.max_row != 12 or sheet.max_column != 12:
        raise ValueError("Unexpected main-table dimensions")
    formulas = [
        cell
        for workbook_sheet in workbook.worksheets
        for row in workbook_sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    if len(formulas) != 121:
        raise ValueError(f"Unexpected formula count: {len(formulas)}")
    japanese_cells = []
    for workbook_sheet in workbook.worksheets:
        for row in workbook_sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and JAPANESE_PATTERN.search(cell.value):
                    japanese_cells.append(f"{workbook_sheet.title}!{cell.coordinate}")
    if japanese_cells:
        raise ValueError(f"Japanese text remains in workbook: {japanese_cells[:5]}")
    if exact_location_total(audit) != 36 or fallback_location_total(audit) != 17:
        raise ValueError("Location totals changed")


def main() -> None:
    audit = build_audit()
    write_workbook(audit)
    validate_output(audit)
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print("Main table: 11 audit components x 12 columns")
    print("Residential boundary differences: Kumamoto +336; Yatsushiro +22")
    print("Calibration differences: 0 for all four city-component checks")
    print("Sex reconciliation: 88/88; location linkage: 36 exact, 17 fallback")


if __name__ == "__main__":
    main()
