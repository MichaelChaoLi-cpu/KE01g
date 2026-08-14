#!/usr/bin/env python3
"""Shelter-Level Equity Demand Scenario Estimates.

Plan: Present an English-only, city-balanced main table containing the ten
highest Base functional-support-demand shelters per city. Preserve the complete
53-shelter, three-scenario evidence in the workbook appendix.
Framework: Section 5's shelter-level scenario contrast; Section 6's catchment
shares, synthetic demand equations, and accessible-parity sensitivity; Section
7's complete shelter equity-demand workflow. Scenario values are planning
estimates, not observed shelter composition or accessible-toilet compliance.
"""

from __future__ import annotations

from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/processed/shelter_equity_scenarios_preprocessed.parquet"
OUTPUT = (
    ROOT
    / "data/results/tables/"
    "Table_shelter_level_equity_demand_scenario_estimates.xlsx"
)

SCENARIOS = ("low", "base", "high")
JAPANESE_PATTERN = re.compile(r"[\u3040-\u30ff\u3400-\u9fff]")
MUNICIPALITY_LABELS = {"八代市": "Yatsushiro City", "熊本市": "Kumamoto City"}
MUNICIPALITY_PREFIXES = {"八代市": "Y", "熊本市": "K"}
LOCATION_LABELS = {
    "exact shelter master match": "Exact shelter match",
    "district anchor fallback": "District-anchor fallback",
}

ENGLISH_SHELTER_NAMES = {
    "Y01": "Daiyo Community Center",
    "Y02": "First Junior High School",
    "Y03": "Yatsushiro Community Center",
    "Y04": "Yatsushiro Elementary School",
    "Y05": "Yatsushiro Toyooka Chiken Arena (General Gymnasium)",
    "Y06": "Otago Community Center",
    "Y07": "Second Junior High School",
    "Y08": "Uyanagi Community Center",
    "Y09": "Uyanagi Elementary School",
    "Y10": "Mugishima Community Center",
    "Y11": "Mugishima Elementary School",
    "Y12": "Matsutaka Community Center",
    "Y13": "Matsutaka Elementary School",
    "Y14": "Fourth Junior High School",
    "Y15": "Yachiwa Community Center",
    "Y16": "Koda Community Center",
    "Y17": "Fifth Junior High School",
    "Y18": "Yatsushiro City Health Center",
    "Y19": "Kongo Community Center",
    "Y20": "Sixth Junior High School",
    "Y21": "Gunchiku Community Center",
    "Y22": "Showa Community Center",
    "Y23": "Miyaji Community Center",
    "Y24": "Miyaji Higashi Community Center",
    "Y25": "Eighth Junior High School",
    "Y26": "Ryuho Community Center",
    "Y27": "Hinagu Community Center",
    "Y28": "Hinagu Elementary School",
    "Y29": "Futami Community Center",
    "Y30": "Sakamoto Community Center",
    "Y31": "Sencho Community Center",
    "Y32": "Kagami Community Center",
    "Y33": "Kagami Elementary School",
    "Y34": "Kagami Junior High School",
    "Y35": "Bunsei Elementary School",
    "Y36": "Arisa Elementary School",
    "Y37": "Toyo Community Center and Toyo Sports Center",
    "Y38": "Yatsushiro Special Needs School",
    "K01": "Oe Community Room and Public Hall",
    "K02": "Akitsu Community Development Center and Public Hall",
    "K03": "Hanazono Community Development Center and Public Hall",
    "K04": "Tatsuda Community Development Center and Public Hall",
    "K07": "Tenmei Community Development Center and Public Hall",
    "K08": "Kawashiri Elementary School",
    "K09": "Nanbu Community Development Center and Public Hall",
    "K12": "Aspal Tomiai (Tomiai Public Hall)",
    "K13": "Koda Community Development Center and Public Hall",
    "K14": "Rikigo Elementary School",
    "K15": "Jonan Senior Welfare Center",
    "K16": "Hinokimi Cultural Center",
    "K17": "Shimomashiki Jonan Junior High School",
    "K18": "Sugiue Elementary School",
    "K19": "Toyoda Elementary School",
}

AREA_LABELS = {
    "代陽": "Daiyo",
    "八代": "Yatsushiro",
    "太田郷": "Otago",
    "植柳": "Uyanagi",
    "麦島": "Mugishima",
    "松高": "Matsutaka",
    "八千把": "Yachiwa",
    "高田": "Koda",
    "金剛": "Kongo",
    "郡築": "Gunchiku",
    "昭和": "Showa",
    "宮地": "Miyaji",
    "龍峯": "Ryuho",
    "日奈久": "Hinagu",
    "二見": "Futami",
    "坂本": "Sakamoto",
    "千丁": "Sencho",
    "鏡": "Kagami",
    "東陽": "Toyo",
    "中央区": "Chuo Ward",
    "東区": "Higashi Ward",
    "西区": "Nishi Ward",
    "北区": "Kita Ward",
    "南区": "Minami Ward",
}

HEADERS = (
    "Stable Shelter ID",
    "Municipality",
    "Shelter Name",
    "Area",
    "Evacuees",
    "Estimated Female Evacuees",
    "Estimated Functional-Support Evacuees: Low",
    "Estimated Functional-Support Evacuees: Base",
    "Estimated Functional-Support Evacuees: High",
    "Estimated Female Functional-Support Evacuees: Low",
    "Estimated Female Functional-Support Evacuees: Base",
    "Estimated Female Functional-Support Evacuees: High",
    "Initial Accessible-Unit Parity Screen (Low / Base / High)",
    "Prolonged Accessible-Unit Parity Screen (Low / Base / High)",
    "Location Resolution",
)


def stable_id(frame: pd.DataFrame) -> pd.Series:
    prefixes = frame["Municipality"].map(MUNICIPALITY_PREFIXES)
    if prefixes.isna().any():
        raise ValueError("Unexpected municipality in matched shelter sample")
    numbers = frame["Shelter Number"].astype(int).astype(str).str.zfill(2)
    return prefixes + numbers


def validate_source(frame: pd.DataFrame) -> pd.DataFrame:
    required = {
        "Municipality",
        "Shelter Number",
        "Shelter Name",
        "District",
        "Ward",
        "Scenario",
        "Evacuees",
        "Estimated Female Evacuees",
        "Estimated Functional Support Evacuees",
        "Estimated Female Functional Support Evacuees",
        "Initial Accessible Unit Parity Screen",
        "Prolonged Accessible Unit Parity Screen",
        "Location Resolution",
        "Evidence Status",
    }
    missing = required.difference(frame.columns)
    if missing:
        raise KeyError(f"Missing required columns: {sorted(missing)}")
    frame = frame.loc[
        frame["Municipality"].isin(MUNICIPALITY_LABELS)
        & frame["Scenario"].astype("string").isin(SCENARIOS)
    ].copy()
    if len(frame) != 159:
        raise ValueError("Expected 53 shelters under each of three scenarios")
    if frame.duplicated(["Municipality", "Shelter Number", "Scenario"]).any():
        raise ValueError("Shelter-scenario keys must be unique")
    if frame[[
        "Municipality",
        "Shelter Number",
        "Shelter Name",
        "Scenario",
        "Evacuees",
        "Estimated Female Evacuees",
        "Estimated Functional Support Evacuees",
        "Estimated Female Functional Support Evacuees",
        "Initial Accessible Unit Parity Screen",
        "Prolonged Accessible Unit Parity Screen",
        "Location Resolution",
        "Evidence Status",
    ]].isna().any().any():
        raise ValueError("Required shelter equity-demand inputs must be complete")
    if set(frame["Scenario"].astype(str)) != set(SCENARIOS):
        raise ValueError("Expected exactly low, base, and high scenarios")
    if set(frame["Evidence Status"].astype(str)) != {
        "scenario estimate, not observed shelter composition"
    }:
        raise ValueError("Unexpected shelter scenario evidence status")
    if not set(frame["Location Resolution"].astype(str)).issubset(LOCATION_LABELS):
        raise ValueError("Unexpected location-resolution category")

    frame["Stable Shelter ID"] = stable_id(frame)
    for shelter_id, rows in frame.groupby("Stable Shelter ID", sort=False):
        if len(rows) != 3:
            raise ValueError(f"Expected three scenario rows for {shelter_id}")
        for column in (
            "Municipality",
            "Shelter Name",
            "Evacuees",
            "Estimated Female Evacuees",
            "Location Resolution",
        ):
            if rows[column].nunique(dropna=False) != 1:
                raise ValueError(f"Observed field varies across scenarios: {shelter_id}, {column}")
        ordered = rows.assign(
            _scenario_order=pd.Categorical(
                rows["Scenario"].astype(str), categories=SCENARIOS, ordered=True
            )
        ).sort_values("_scenario_order")
        for column in (
            "Estimated Functional Support Evacuees",
            "Estimated Female Functional Support Evacuees",
            "Initial Accessible Unit Parity Screen",
            "Prolonged Accessible Unit Parity Screen",
        ):
            values = ordered[column].to_numpy(dtype=float)
            if not np.all(values[:-1] <= values[1:]):
                raise ValueError(f"Scenario values are not monotonic: {shelter_id}, {column}")
    base = frame.loc[frame["Scenario"].astype(str).eq("base")]
    city_checks = base.groupby("Municipality").agg(
        Shelters=("Stable Shelter ID", "size"), Evacuees=("Evacuees", "sum")
    )
    expected = {"八代市": (38, 1852), "熊本市": (15, 280)}
    for municipality, (shelters, evacuees) in expected.items():
        row = city_checks.loc[municipality]
        if int(row["Shelters"]) != shelters or int(row["Evacuees"]) != evacuees:
            raise ValueError(f"Matched city totals changed for {municipality}")
    return frame


def prepare_inputs(frame: pd.DataFrame) -> pd.DataFrame:
    scenario_order = pd.Categorical(
        frame["Scenario"].astype(str), categories=SCENARIOS, ordered=True
    )
    municipality_order = pd.Categorical(
        frame["Municipality"], categories=["八代市", "熊本市"], ordered=True
    )
    prepared = frame.assign(
        _municipality_order=municipality_order, _scenario_order=scenario_order
    ).sort_values(
        ["_municipality_order", "Shelter Number", "_scenario_order"], kind="stable"
    )
    return prepared.reset_index(drop=True)


def write_workbook(inputs_frame: pd.DataFrame) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Equity Demand Summary"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "E2"
    sheet.sheet_view.zoomScale = 70
    sheet.append(list(HEADERS))

    input_rows: dict[tuple[str, str], int] = {}
    for excel_row, (_, row) in enumerate(inputs_frame.iterrows(), start=2):
        input_rows[
            (str(row["Stable Shelter ID"]), str(row["Scenario"]))
        ] = excel_row
    scenario_values = inputs_frame.assign(
        Scenario=inputs_frame["Scenario"].astype(str)
    ).set_index(["Stable Shelter ID", "Scenario"])

    full_shelter_order = (
        inputs_frame.loc[inputs_frame["Scenario"].astype(str).eq("base")]
        .sort_values(["_municipality_order", "Shelter Number"], kind="stable")
        .reset_index(drop=True)
    )
    selected_frames: list[pd.DataFrame] = []
    for municipality in ("八代市", "熊本市"):
        city = full_shelter_order.loc[
            full_shelter_order["Municipality"].eq(municipality)
        ].sort_values(
            [
                "Estimated Functional Support Evacuees",
                "Estimated Female Functional Support Evacuees",
                "Evacuees",
                "Stable Shelter ID",
            ],
            ascending=[False, False, False, True],
            kind="stable",
        )
        selected_frames.append(city.head(10))
    shelter_order = pd.concat(selected_frames, ignore_index=True)
    if len(shelter_order) != 20 or shelter_order.groupby("Municipality").size().to_dict() != {
        "八代市": 10,
        "熊本市": 10,
    }:
        raise ValueError("Expected the ten highest Base-demand shelters in each city")
    for _, row in shelter_order.iterrows():
        shelter_id = str(row["Stable Shelter ID"])
        low = scenario_values.loc[(shelter_id, "low")]
        base = scenario_values.loc[(shelter_id, "base")]
        high = scenario_values.loc[(shelter_id, "high")]
        area_value = row["District"] if pd.notna(row["District"]) else row["Ward"]
        sheet.append(
            [
                shelter_id,
                MUNICIPALITY_LABELS[str(row["Municipality"])],
                ENGLISH_SHELTER_NAMES[shelter_id],
                AREA_LABELS[str(area_value)],
                int(row["Evacuees"]),
                float(base["Estimated Female Evacuees"]),
                float(low["Estimated Functional Support Evacuees"]),
                float(base["Estimated Functional Support Evacuees"]),
                float(high["Estimated Functional Support Evacuees"]),
                float(low["Estimated Female Functional Support Evacuees"]),
                float(base["Estimated Female Functional Support Evacuees"]),
                float(high["Estimated Female Functional Support Evacuees"]),
                " / ".join(
                    str(int(value))
                    for value in (
                        low["Initial Accessible Unit Parity Screen"],
                        base["Initial Accessible Unit Parity Screen"],
                        high["Initial Accessible Unit Parity Screen"],
                    )
                ),
                " / ".join(
                    str(int(value))
                    for value in (
                        low["Prolonged Accessible Unit Parity Screen"],
                        base["Prolonged Accessible Unit Parity Screen"],
                        high["Prolonged Accessible Unit Parity Screen"],
                    )
                ),
                LOCATION_LABELS[str(row["Location Resolution"])],
            ]
        )

    last_row = sheet.max_row
    table_range = f"A1:O{last_row}"
    excel_table = Table(displayName="ShelterEquityDemandSummary", ref=table_range)
    excel_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(excel_table)

    header_fill = PatternFill("solid", fgColor="27445C")
    header_font = Font(name="Aptos", size=8.5, bold=True, color="FFFFFF")
    body_font = Font(name="Aptos", size=7.8, color="263746")
    thin_gray = Side(style="thin", color="D8DEE3")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="9EB1C0"))
    sheet.row_dimensions[1].height = 78
    for row in sheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=15):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
            cell.border = Border(bottom=thin_gray)
        row[2].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        for index in range(4, 14):
            row[index].alignment = Alignment(horizontal="right", vertical="center")
        row[12].alignment = Alignment(horizontal="center", vertical="center")
        row[13].alignment = Alignment(horizontal="center", vertical="center")
        row[14].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 38
    for column in ("E",):
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0"
    for column in ("F", "G", "H", "I", "J", "K", "L"):
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0.0"

    sheet.conditional_formatting.add(
        f"O2:O{last_row}",
        FormulaRule(
            formula=['LEFT($O2,5)="Exact"'],
            fill=PatternFill("solid", fgColor="DCEDE5"),
        ),
    )
    sheet.conditional_formatting.add(
        f"O2:O{last_row}",
        FormulaRule(
            formula=['LEFT($O2,8)="District"'],
            fill=PatternFill("solid", fgColor="FFF0CC"),
        ),
    )
    sheet.conditional_formatting.add(
        f"A2:A{last_row}",
        FormulaRule(
            formula=['LEFT($A2,1)="Y"'],
            fill=PatternFill("solid", fgColor="E6F0F7"),
        ),
    )

    widths = {
        "A": 11,
        "B": 15,
        "C": 35,
        "D": 14,
        "E": 9,
        "F": 14,
        "G": 14,
        "H": 14,
        "I": 14,
        "J": 15,
        "K": 15,
        "L": 15,
        "M": 18,
        "N": 19,
        "O": 20,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell_ref, text in {
        "F1": "Observed evacuees multiplied by the residential catchment female share; not observed female shelter attendance.",
        "G1": "Low assumes complete overlap between allocated disability and long-term-care components.",
        "H1": "Base assumes half overlap between allocated disability and long-term-care components.",
        "I1": "High assumes no overlap between allocated disability and long-term-care components.",
        "M1": "Low / Base / High scenario order. Applies the 1:50 ratio to estimated functional-support demand; not an official accessible-toilet threshold.",
        "N1": "Low / Base / High scenario order. Applies the 1:20 ratio to estimated functional-support demand; not an official accessible-toilet threshold.",
        "O1": "District-anchor fallback supports district screening only, not precise shelter-catchment interpretation.",
    }.items():
        sheet[cell_ref].comment = Comment(text, "Mike Li")
    sheet.auto_filter.ref = table_range
    sheet.print_area = table_range
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.14
    sheet.page_margins.right = 0.14
    sheet.page_margins.top = 0.25
    sheet.page_margins.bottom = 0.25

    inputs = workbook.create_sheet("Scenario Inputs")
    inputs.sheet_view.showGridLines = False
    input_headers = (
        "Stable Shelter ID",
        "Municipality",
        "Shelter Number",
        "Shelter Name",
        "Scenario",
        "Evacuees",
        "Catchment Female Share",
        "Estimated Female Evacuees",
        "Estimated Functional Support Evacuees",
        "Estimated Female Functional Support Evacuees",
        "Initial Accessible Unit Parity Screen",
        "Prolonged Accessible Unit Parity Screen",
        "Location Resolution",
        "Evidence Status",
    )
    inputs.append(list(input_headers))
    for _, row in inputs_frame.iterrows():
        inputs.append(
            [
                str(row["Stable Shelter ID"]),
                MUNICIPALITY_LABELS[str(row["Municipality"])],
                int(row["Shelter Number"]),
                ENGLISH_SHELTER_NAMES[str(row["Stable Shelter ID"])],
                str(row["Scenario"]),
                int(row["Evacuees"]),
                float(row["Catchment Female Share"]),
                float(row["Estimated Female Evacuees"]),
                float(row["Estimated Functional Support Evacuees"]),
                float(row["Estimated Female Functional Support Evacuees"]),
                int(row["Initial Accessible Unit Parity Screen"]),
                int(row["Prolonged Accessible Unit Parity Screen"]),
                str(row["Location Resolution"]),
                str(row["Evidence Status"]),
            ]
        )
    for cell in inputs[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    inputs.row_dimensions[1].height = 58
    for row in inputs.iter_rows(min_row=2, max_row=inputs.max_row, min_col=1, max_col=14):
        for cell in row:
            cell.font = Font(name="Aptos", size=8, color="263746")
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.border = Border(bottom=thin_gray)
        for index in (0, 1, 3, 4, 12, 13):
            row[index].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for column in ("G",):
        for cell in inputs[column][1:]:
            cell.number_format = "0.0%"
    for column in ("H", "I", "J"):
        for cell in inputs[column][1:]:
            cell.number_format = "#,##0.0"
    input_widths = {
        "A": 11,
        "B": 15,
        "C": 10,
        "D": 26,
        "E": 10,
        "F": 10,
        "G": 14,
        "H": 15,
        "I": 17,
        "J": 18,
        "K": 16,
        "L": 17,
        "M": 23,
        "N": 37,
    }
    for column, width in input_widths.items():
        inputs.column_dimensions[column].width = width
    inputs.freeze_panes = "F2"
    inputs.sheet_view.zoomScale = 70
    inputs.print_area = f"A1:N{inputs.max_row}"
    inputs.print_title_rows = "1:1"
    inputs.page_setup.orientation = "landscape"
    inputs.page_setup.paperSize = inputs.PAPERSIZE_A3
    inputs.page_setup.fitToWidth = 1
    inputs.page_setup.fitToHeight = 0
    inputs.sheet_properties.pageSetUpPr.fitToPage = True

    notes = workbook.create_sheet("Notes")
    notes.sheet_view.showGridLines = False
    notes["A1"] = "Scenario definitions, evidence boundaries, and sources"
    notes["A1"].font = Font(name="Aptos Display", size=15, bold=True, color="27445C")
    notes.append([])
    notes.append(["Item", "Definition or boundary", "Source URL"])
    note_rows = (
        (
            "Main-table selection",
            "The main sheet retains the ten shelters with the highest Base functional-support demand in each city. The Scenario Inputs appendix preserves all 53 shelters under Low, Base, and High.",
            "",
        ),
        (
            "Observed field",
            "Evacuees are observed in matched city shelter tables at 12:00 JST on 12 August 2026.",
            "",
        ),
        (
            "Estimated female demand",
            "Observed occupancy multiplied by the 2020 residential catchment female share; not observed shelter composition.",
            "https://www.e-stat.go.jp/gis/statmap-search?page=1&type=1&toukeiCode=00200521",
        ),
        (
            "Functional-support scenarios",
            "Low, Base, and High assume complete, half, and no overlap between calibrated physical-disability and long-term-care components. They are planning assumptions, not confidence intervals.",
            "",
        ),
        (
            "Accessible-unit parity screen",
            "The initial and prolonged screens apply the ordinary 1:50 and 1:20 ratios to estimated functional-support demand. They are not official accessible-toilet compliance requirements.",
            "https://www.bousai.go.jp/taisaku/hinanjo/pdf/2412hinanjo_toilet_guideline.pdf",
        ),
        (
            "Exact shelter match",
            "Coordinates match the shelter master and support shelter-level scenario interpretation, subject to all other evidence limits.",
            "",
        ),
        (
            "District-anchor fallback",
            "Seventeen Yatsushiro shelter locations use district anchors and support district screening only, not precise shelter catchments.",
            "",
        ),
        (
            "Interpretation boundary",
            "The table does not observe sex, disability, care need, toilet accessibility, fixed stalls, or functional toilet capacity among shelter users.",
            "",
        ),
    )
    for values in note_rows:
        notes.append(values)
    for cell in notes[3]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
    notes.row_dimensions[3].height = 24
    for row in notes.iter_rows(min_row=4, max_row=notes.max_row, min_col=1, max_col=3):
        row[0].font = Font(name="Aptos", size=9, bold=True, color="263746")
        row[1].font = Font(name="Aptos", size=8.5, color="263746")
        row[2].font = Font(name="Aptos", size=8, color="0563C1", underline="single")
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
        notes.row_dimensions[row[0].row].height = 42
    notes.column_dimensions["A"].width = 28
    notes.column_dimensions["B"].width = 100
    notes.column_dimensions["C"].width = 78
    notes.freeze_panes = "A4"
    notes.sheet_view.zoomScale = 85
    notes.print_area = f"A1:C{notes.max_row}"
    notes.page_setup.orientation = "landscape"
    notes.page_setup.paperSize = notes.PAPERSIZE_A4
    notes.page_setup.fitToWidth = 1
    notes.page_setup.fitToHeight = 1
    notes.sheet_properties.pageSetUpPr.fitToPage = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)


def verify_workbook() -> None:
    workbook = load_workbook(OUTPUT, data_only=False, read_only=False)
    if workbook.sheetnames != ["Equity Demand Summary", "Scenario Inputs", "Notes"]:
        raise ValueError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    sheet = workbook["Equity Demand Summary"]
    if sheet.max_row != 21 or sheet.max_column != 15:
        raise ValueError(f"Unexpected summary dimensions: {sheet.max_row} x {sheet.max_column}")
    if tuple(cell.value for cell in sheet[1]) != HEADERS:
        raise ValueError("Summary headers do not match the planned 15-column table")
    if any(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for row in sheet.iter_rows()
        for cell in row
    ):
        raise ValueError("The article-facing Equity Demand Summary must contain static values")
    if {
        city: sum(sheet.cell(row=row, column=2).value == city for row in range(2, 22))
        for city in MUNICIPALITY_LABELS.values()
    } != {"Yatsushiro City": 10, "Kumamoto City": 10}:
        raise ValueError("Main table is not balanced at ten shelters per city")
    inputs = workbook["Scenario Inputs"]
    if inputs.max_row != 160 or inputs.max_column != 14:
        raise ValueError("Scenario Inputs must contain 159 rows and 14 columns")
    japanese_cells = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and JAPANESE_PATTERN.search(cell.value):
                    japanese_cells.append(f"{worksheet.title}!{cell.coordinate}")
    if japanese_cells:
        raise ValueError(f"Japanese text remains in workbook: {japanese_cells[:5]}")
    workbook.close()


def main() -> None:
    source = validate_source(pd.read_parquet(SOURCE))
    inputs = prepare_inputs(source)
    write_workbook(inputs)
    verify_workbook()
    base = source.loc[source["Scenario"].astype(str).eq("base")]
    exact = int(base["Location Resolution"].astype(str).eq("exact shelter master match").sum())
    fallback = int(base["Location Resolution"].astype(str).eq("district anchor fallback").sum())
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print("Main summary: 20 shelters x 15 columns (10 per city)")
    print("Appendix: complete 53-shelter sample under 3 scenarios (159 rows)")
    print(f"Location resolution: {exact} exact shelter matches; {fallback} district-anchor fallbacks")


if __name__ == "__main__":
    main()
