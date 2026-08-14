#!/usr/bin/env python3
"""Water Status and Equity-Weighted Pressure Summary.

Plan: Summarize observed demand, women and functional-support scenario demand,
reported deployment, and temporary-toilet-only pressure for each Yatsushiro
water-service category under all three overlap scenarios.
Framework: Section 5 water-status stratification; Section 6 water-severity,
grouped-summary, prolonged requirement, and screening-shortfall formulas;
Section 7 grouped pressure workflow. Water status is an operability proxy and
the shortfall is not a functional-capacity gap.
"""

from __future__ import annotations

from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/processed/shelter_equity_scenarios_preprocessed.parquet"
OUTPUT = (
    ROOT
    / "data/results/tables/"
    "Table_water_status_and_equity_weighted_pressure_summary.xlsx"
)

SCENARIO_ORDER = ("Low", "Base", "High")
WATER_ORDER = ("Available", "Partially available", "Unavailable")
WATER_LABELS = {
    "〇": "Available",
    "○": "Available",
    "△": "Partially available",
    "×": "Unavailable",
}
WATER_FILLS = {
    "Available": "DCEDE5",
    "Partially available": "FFF0CC",
    "Unavailable": "F5D7D7",
}
JAPANESE_PATTERN = re.compile(r"[\u3040-\u30ff\u3400-\u9fff]")
HEADERS = (
    "Water Status",
    "Scenario",
    "Shelters",
    "Evacuees",
    "Estimated Female Evacuees",
    "Estimated Functional-Support Evacuees",
    "Estimated Female Functional-Support Evacuees",
    "Temporary Toilets Installed",
    "Toilet Cars",
    "Prolonged Requirement",
    "Temporary-Toilet-Only Shortfall",
    "Shelters with Positive Shortfall",
)


def prepare_source(frame: pd.DataFrame) -> pd.DataFrame:
    required = {
        "Municipality",
        "Shelter Number",
        "Scenario",
        "Evacuees",
        "Estimated Female Evacuees",
        "Estimated Functional Support Evacuees",
        "Estimated Female Functional Support Evacuees",
        "Water Status",
        "Temporary Toilets Installed",
        "Toilet Cars",
    }
    missing = required.difference(frame.columns)
    if missing:
        raise KeyError(f"Missing required columns: {sorted(missing)}")
    yatsushiro = frame.loc[
        frame["Municipality"].eq("八代市")
        & frame["Scenario"].astype("string").isin(["low", "base", "high"])
    ].copy()
    if len(yatsushiro) != 114:
        raise ValueError("Expected 38 Yatsushiro shelters under three scenarios")
    if yatsushiro.duplicated(["Shelter Number", "Scenario"]).any():
        raise ValueError("Shelter-scenario keys must be unique")
    completeness_columns = list(required.difference({"Municipality"}))
    if yatsushiro[completeness_columns].isna().any().any():
        raise ValueError("Yatsushiro water-pressure inputs must be complete")
    yatsushiro["Water Group"] = yatsushiro["Water Status"].astype(str).map(WATER_LABELS)
    if yatsushiro["Water Group"].isna().any():
        raise ValueError("Unrecognized water-status value")
    counts = (
        yatsushiro.loc[yatsushiro["Scenario"].astype(str).eq("base"), "Water Group"]
        .value_counts()
        .to_dict()
    )
    expected_counts = {"Available": 19, "Partially available": 6, "Unavailable": 13}
    if counts != expected_counts:
        raise ValueError(f"Unexpected water-status counts: {counts}")
    yatsushiro["Scenario English"] = yatsushiro["Scenario"].astype(str).str.title()
    yatsushiro["Stable Shelter ID"] = (
        "Y" + yatsushiro["Shelter Number"].astype(int).astype(str).str.zfill(2)
    )
    yatsushiro["Prolonged Requirement"] = np.ceil(yatsushiro["Evacuees"] / 20).astype(int)
    yatsushiro["Temporary-Toilet-Only Shortfall"] = (
        yatsushiro["Prolonged Requirement"]
        - yatsushiro["Temporary Toilets Installed"]
    ).clip(lower=0).astype(int)
    shortfall_total = yatsushiro.loc[
        yatsushiro["Scenario"].astype(str).eq("base"),
        "Temporary-Toilet-Only Shortfall",
    ].sum()
    if int(shortfall_total) != 31:
        raise ValueError("Unexpected Yatsushiro screening-shortfall total")
    water_order = pd.Categorical(
        yatsushiro["Water Group"], categories=WATER_ORDER, ordered=True
    )
    scenario_order = pd.Categorical(
        yatsushiro["Scenario English"], categories=SCENARIO_ORDER, ordered=True
    )
    return (
        yatsushiro.assign(_water_order=water_order, _scenario_order=scenario_order)
        .sort_values(["_water_order", "_scenario_order", "Shelter Number"], kind="stable")
        .reset_index(drop=True)
    )


def expected_summary(source: pd.DataFrame) -> pd.DataFrame:
    summary = (
        source.groupby(["Water Group", "Scenario English"], observed=True, sort=False)
        .agg(
            Shelters=("Stable Shelter ID", "size"),
            Evacuees=("Evacuees", "sum"),
            Estimated_Female_Evacuees=("Estimated Female Evacuees", "sum"),
            Estimated_Functional_Support_Evacuees=(
                "Estimated Functional Support Evacuees",
                "sum",
            ),
            Estimated_Female_Functional_Support_Evacuees=(
                "Estimated Female Functional Support Evacuees",
                "sum",
            ),
            Temporary_Toilets_Installed=("Temporary Toilets Installed", "sum"),
            Toilet_Cars=("Toilet Cars", "sum"),
            Prolonged_Requirement=("Prolonged Requirement", "sum"),
            Temporary_Toilet_Only_Shortfall=(
                "Temporary-Toilet-Only Shortfall",
                "sum",
            ),
            Shelters_with_Positive_Shortfall=(
                "Temporary-Toilet-Only Shortfall",
                lambda values: int(values.gt(0).sum()),
            ),
        )
        .reset_index()
    )
    if len(summary) != 9:
        raise ValueError("Expected three water groups by three scenarios")
    return summary


def write_workbook(source: pd.DataFrame, summary: pd.DataFrame) -> None:
    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    sheet = workbook.active
    sheet.title = "Water Pressure Summary"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C2"
    sheet.sheet_view.zoomScale = 78
    sheet.append(list(HEADERS))

    source_last_row = len(source) + 1
    source_columns = {
        "Shelters": "A",
        "Evacuees": "D",
        "Estimated Female Evacuees": "E",
        "Estimated Functional-Support Evacuees": "F",
        "Estimated Female Functional-Support Evacuees": "G",
        "Temporary Toilets Installed": "H",
        "Toilet Cars": "I",
        "Prolonged Requirement": "J",
        "Temporary-Toilet-Only Shortfall": "K",
    }
    output_rows = [(water, scenario) for water in WATER_ORDER for scenario in SCENARIO_ORDER]
    for row_number, (water, scenario) in enumerate(output_rows, start=2):
        criteria = (
            f"'Scenario Inputs'!$B$2:$B${source_last_row},$A{row_number},"
            f"'Scenario Inputs'!$C$2:$C${source_last_row},$B{row_number}"
        )
        values: list[object] = [water, scenario]
        values.append(f'=COUNTIFS({criteria})')
        for header in HEADERS[3:11]:
            column = source_columns[header]
            values.append(
                f"=SUMIFS('Scenario Inputs'!${column}$2:${column}${source_last_row},{criteria})"
            )
        values.append(
            f'=COUNTIFS({criteria},\'Scenario Inputs\'!$K$2:$K${source_last_row},">0")'
        )
        sheet.append(values)

    last_row = sheet.max_row
    excel_table = Table(displayName="WaterEquityPressureSummary", ref=f"A1:L{last_row}")
    excel_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(excel_table)

    navy = "27445C"
    text = "263746"
    border = Side(style="thin", color="D8DEE3")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=8.5, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="9EB1C0"))
    sheet.row_dimensions[1].height = 68
    for row in sheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=12):
        for cell in row:
            cell.font = Font(name="Aptos", size=8.5, color=text)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=border)
        for index in range(2, 12):
            row[index].alignment = Alignment(horizontal="right", vertical="center")
        sheet.row_dimensions[row[0].row].height = 31
        row[0].fill = PatternFill("solid", fgColor=WATER_FILLS[str(row[0].value)])
        if row[1].value == "Base":
            row[1].font = Font(name="Aptos", size=8.5, bold=True, color="27445C")
    for column in ("C", "D", "H", "I", "J", "K", "L"):
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0"
    for column in ("E", "F", "G"):
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0.0"
    sheet.conditional_formatting.add(
        f"K2:K{last_row}",
        ColorScaleRule(
            start_type="min",
            start_color="E4EFEA",
            mid_type="percentile",
            mid_value=50,
            mid_color="F4DEB6",
            end_type="max",
            end_color="D99A8E",
        ),
    )
    widths = {
        "A": 19,
        "B": 10,
        "C": 10,
        "D": 11,
        "E": 17,
        "F": 20,
        "G": 23,
        "H": 18,
        "I": 10,
        "J": 16,
        "K": 21,
        "L": 20,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet["K1"].comment = Comment(
        "This screen equals max(0, prolonged benchmark requirement minus reported temporary toilets), summed by water group. It excludes fixed stalls and toilet cars.",
        "OpenAI",
    )
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = f"A1:L{last_row}"

    inputs = workbook.create_sheet("Scenario Inputs")
    inputs.sheet_view.showGridLines = False
    input_headers = (
        "Stable Shelter ID",
        "Water Status",
        "Scenario",
        "Evacuees",
        "Estimated Female Evacuees",
        "Estimated Functional-Support Evacuees",
        "Estimated Female Functional-Support Evacuees",
        "Temporary Toilets Installed",
        "Toilet Cars",
        "Prolonged Requirement",
        "Temporary-Toilet-Only Shortfall",
    )
    inputs.append(input_headers)
    for row_number, (_, row) in enumerate(source.iterrows(), start=2):
        inputs.append(
            [
                row["Stable Shelter ID"],
                row["Water Group"],
                row["Scenario English"],
                int(row["Evacuees"]),
                float(row["Estimated Female Evacuees"]),
                float(row["Estimated Functional Support Evacuees"]),
                float(row["Estimated Female Functional Support Evacuees"]),
                int(row["Temporary Toilets Installed"]),
                int(row["Toilet Cars"]),
                f"=ROUNDUP(D{row_number}/20,0)",
                f"=MAX(0,J{row_number}-H{row_number})",
            ]
        )
    for cell in inputs[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=8.5, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    inputs.row_dimensions[1].height = 60
    for row in inputs.iter_rows(min_row=2, max_row=inputs.max_row):
        for cell in row:
            cell.font = Font(name="Aptos", size=8, color=text)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = Border(bottom=border)
        for index in range(3, 11):
            row[index].alignment = Alignment(horizontal="right", vertical="center")
    for column, width in {
        "A": 17,
        "B": 19,
        "C": 10,
        "D": 11,
        "E": 17,
        "F": 21,
        "G": 24,
        "H": 20,
        "I": 11,
        "J": 18,
        "K": 22,
    }.items():
        inputs.column_dimensions[column].width = width
    inputs.freeze_panes = "D2"
    inputs.auto_filter.ref = f"A1:K{inputs.max_row}"

    notes = workbook.create_sheet("Notes")
    notes.sheet_view.showGridLines = False
    notes.column_dimensions["A"].width = 32
    notes.column_dimensions["B"].width = 108
    notes.append(["Field", "Definition"])
    note_rows = [
        ("Scope", "Yatsushiro City's 38 matched shelters at the common observation, shown under Low, Base, and High women–functional-support overlap scenarios."),
        ("Water status", "Reported source category used as an operability proxy: available, partially available, or unavailable. It does not verify wastewater service or individual toilet functionality."),
        ("Scenario demand", "Female, functional-support, and female functional-support values are synthetic expectations, not observed evacuee subgroup counts."),
        ("Temporary-toilet-only shortfall", "The positive difference between each shelter's prolonged benchmark requirement and reported temporary toilets, summed within each group. Fixed stalls and toilet cars are excluded."),
        ("Interpretation", "The table identifies co-location for field verification. It is not a functional-capacity deficit, accessible-toilet compliance result, or causal effect of water status."),
    ]
    for row in note_rows:
        notes.append(row)
    for cell in notes[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
    for row in notes.iter_rows(min_row=2, max_row=notes.max_row):
        row[0].font = Font(name="Aptos", size=9, bold=True, color=text)
        row[1].font = Font(name="Aptos", size=9, color=text)
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = Border(bottom=border)
        notes.row_dimensions[row[0].row].height = 46
    notes.freeze_panes = "A2"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)


def validate_output(summary: pd.DataFrame) -> None:
    workbook = load_workbook(OUTPUT, data_only=False)
    if workbook.sheetnames != ["Water Pressure Summary", "Scenario Inputs", "Notes"]:
        raise ValueError("Unexpected workbook sheet structure")
    sheet = workbook["Water Pressure Summary"]
    if sheet.max_row != 10 or sheet.max_column != 12:
        raise ValueError("Unexpected main-table dimensions")
    formula_count = sum(
        1
        for workbook_sheet in workbook.worksheets
        for row in workbook_sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )
    if formula_count != 318:
        raise ValueError(f"Unexpected formula count: {formula_count}")
    japanese_cells = []
    for workbook_sheet in workbook.worksheets:
        for row in workbook_sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and JAPANESE_PATTERN.search(cell.value):
                    japanese_cells.append(f"{workbook_sheet.title}!{cell.coordinate}")
    if japanese_cells:
        raise ValueError(f"Japanese text remains in workbook: {japanese_cells[:5]}")
    if int(summary["Temporary_Toilet_Only_Shortfall"].sum()) != 93:
        raise ValueError("Shortfall total across three scenarios must equal 3 x 31")
    base = summary.loc[summary["Scenario English"].eq("Base")]
    if int(base["Shelters"].sum()) != 38 or int(base["Evacuees"].sum()) != 1852:
        raise ValueError("Base grouped totals do not reproduce the matched observation")


def main() -> None:
    source = prepare_source(pd.read_parquet(SOURCE))
    summary = expected_summary(source)
    write_workbook(source, summary)
    validate_output(summary)
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print("Main table: 9 water-status-scenario rows x 12 columns")
    base = summary.loc[summary["Scenario English"].eq("Base")]
    for _, row in base.iterrows():
        print(
            f"{row['Water Group']}: {int(row['Evacuees'])} evacuees, "
            f"{row['Estimated_Functional_Support_Evacuees']:.1f} base FS demand, "
            f"{int(row['Temporary_Toilet_Only_Shortfall'])} shortfall units"
        )


if __name__ == "__main__":
    main()
