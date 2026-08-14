#!/usr/bin/env python3
"""Aggregate and Sitewise Requirement Comparison.

Plan: Compare pooled-city and summed shelter-specific toilet requirements for
both matched cities under the initial-response and prolonged-stay benchmarks.
Framework: Section 5 spatial-fragmentation contrast; Section 6 benchmark,
pooled, sitewise, and fragmentation-increment equations; Section 7 matched-city
comparison. Requirements are standards-based screens, not functional deficits.
"""

from __future__ import annotations

from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from table_shelter_level_equity_demand_scenario_estimates import (
    MUNICIPALITY_LABELS,
    MUNICIPALITY_PREFIXES,
)


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/processed/shelter_equity_scenarios_preprocessed.parquet"
OUTPUT = (
    ROOT
    / "data/results/tables/"
    "Table_aggregate_and_sitewise_requirement_comparison.xlsx"
)

CITY_ORDER = ("八代市", "熊本市")
BENCHMARKS = (("Initial response (1 per 50)", 50), ("Prolonged stay (1 per 20)", 20))
JAPANESE_PATTERN = re.compile(r"[\u3040-\u30ff\u3400-\u9fff]")
HEADERS = (
    "Municipality",
    "Toilet Benchmark",
    "Open Shelters",
    "Evacuees",
    "Pooled Requirement",
    "Summed Sitewise Requirement",
    "Fragmentation Increment",
    "Fragmentation Increment (%)",
    "Reported Temporary Toilets (Observed / Shelters)",
)


def prepare_source(frame: pd.DataFrame) -> pd.DataFrame:
    required = {
        "Municipality",
        "Shelter Number",
        "Scenario",
        "Evacuees",
        "Temporary Toilets Installed",
    }
    missing = required.difference(frame.columns)
    if missing:
        raise KeyError(f"Missing required columns: {sorted(missing)}")
    base = frame.loc[
        frame["Municipality"].isin(CITY_ORDER)
        & frame["Scenario"].astype("string").eq("base")
    ].copy()
    if len(base) != 53 or base.duplicated(["Municipality", "Shelter Number"]).any():
        raise ValueError("Expected 53 unique shelters in the matched observation")
    if base["Evacuees"].isna().any() or base["Evacuees"].lt(0).any():
        raise ValueError("Evacuees must be complete and nonnegative")
    expected = {"八代市": (38, 1852), "熊本市": (15, 280)}
    for municipality, (shelters, evacuees) in expected.items():
        city = base.loc[base["Municipality"].eq(municipality)]
        if len(city) != shelters or int(city["Evacuees"].sum()) != evacuees:
            raise ValueError(f"Matched sample does not reproduce {municipality} totals")
    base["Stable Shelter ID"] = (
        base["Municipality"].map(MUNICIPALITY_PREFIXES)
        + base["Shelter Number"].astype(int).astype(str).str.zfill(2)
    )
    base["Municipality English"] = base["Municipality"].map(MUNICIPALITY_LABELS)
    base["Initial Requirement"] = np.ceil(base["Evacuees"] / 50).astype(int)
    base["Prolonged Requirement"] = np.ceil(base["Evacuees"] / 20).astype(int)
    return base.sort_values(
        ["Municipality", "Shelter Number"],
        key=lambda values: (
            values.map({"八代市": 0, "熊本市": 1})
            if values.name == "Municipality"
            else values
        ),
        kind="stable",
    ).reset_index(drop=True)


def expected_comparison(base: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for municipality in CITY_ORDER:
        city = base.loc[base["Municipality"].eq(municipality)]
        for benchmark, denominator in BENCHMARKS:
            pooled = int(np.ceil(city["Evacuees"].sum() / denominator))
            sitewise = int(np.ceil(city["Evacuees"] / denominator).sum())
            rows.append(
                {
                    "Municipality": MUNICIPALITY_LABELS[municipality],
                    "Toilet Benchmark": benchmark,
                    "Open Shelters": len(city),
                    "Evacuees": int(city["Evacuees"].sum()),
                    "Pooled Requirement": pooled,
                    "Summed Sitewise Requirement": sitewise,
                    "Fragmentation Increment": sitewise - pooled,
                    "Fragmentation Increment (%)": 100 * (sitewise - pooled) / pooled,
                    "Reported Temporary Toilets": int(
                        city["Temporary Toilets Installed"].sum(min_count=1)
                    ),
                    "Deployment Coverage": int(
                        city["Temporary Toilets Installed"].notna().sum()
                    ),
                }
            )
    result = pd.DataFrame(rows)
    if len(result) != 4:
        raise ValueError("Expected two cities by two benchmarks")
    if (result["Summed Sitewise Requirement"] < result["Pooled Requirement"]).any():
        raise ValueError("Sitewise requirement cannot be lower than pooled requirement")
    return result


def write_workbook(base: pd.DataFrame, comparison: pd.DataFrame) -> None:
    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    sheet = workbook.active
    sheet.title = "Requirement Comparison"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C2"
    sheet.sheet_view.zoomScale = 90
    sheet.append(list(HEADERS))

    for _, row in comparison.iterrows():
        city = str(row["Municipality"])
        benchmark = str(row["Toilet Benchmark"])
        coverage = int(row["Deployment Coverage"])
        reported = (
            f"NR (0/{int(row['Open Shelters'])})"
            if coverage == 0
            else f"{int(row['Reported Temporary Toilets'])} ({coverage}/{int(row['Open Shelters'])})"
        )
        sheet.append(
            [
                city,
                benchmark,
                int(row["Open Shelters"]),
                int(row["Evacuees"]),
                int(row["Pooled Requirement"]),
                int(row["Summed Sitewise Requirement"]),
                int(row["Fragmentation Increment"]),
                float(row["Fragmentation Increment (%)"]) / 100,
                reported,
            ]
        )

    last_row = sheet.max_row
    excel_table = Table(displayName="AggregateSitewiseComparison", ref=f"A1:I{last_row}")
    excel_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(excel_table)

    navy = "27445C"
    text = "263746"
    border = Side(style="thin", color="D8DEE3")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="9EB1C0"))
    sheet.row_dimensions[1].height = 62
    for row in sheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=9):
        for cell in row:
            cell.font = Font(name="Aptos", size=9, color=text)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=border)
        for index in (2, 3, 4, 5, 6, 7):
            row[index].alignment = Alignment(horizontal="right", vertical="center")
        row[8].alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[row[0].row].height = 34
    for column in ("C", "D", "E", "F", "G"):
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0"
    for cell in sheet["H"][1:]:
        cell.number_format = "0.0%"
    sheet.conditional_formatting.add(
        f"H2:H{last_row}",
        ColorScaleRule(
            start_type="min",
            start_color="E7EEF3",
            mid_type="percentile",
            mid_value=50,
            mid_color="F1D9B7",
            end_type="max",
            end_color="D99A8E",
        ),
    )
    widths = {
        "A": 18,
        "B": 25,
        "C": 13,
        "D": 13,
        "E": 16,
        "F": 20,
        "G": 17,
        "H": 19,
        "I": 28,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet["I1"].comment = Comment(
        "The total uses only observed Temporary Toilets Installed values. The parenthetical count is observed shelters divided by all shelters; missing values are not zero.",
        "OpenAI",
    )
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = f"A1:I{last_row}"

    inputs = workbook.create_sheet("Shelter Inputs")
    inputs.sheet_view.showGridLines = False
    input_headers = (
        "Stable Shelter ID",
        "Municipality",
        "Evacuees",
        "Temporary Toilets Installed",
        "Initial Requirement",
        "Prolonged Requirement",
    )
    inputs.append(input_headers)
    for row_number, (_, row) in enumerate(base.iterrows(), start=2):
        temp_value = row["Temporary Toilets Installed"]
        inputs.append(
            [
                row["Stable Shelter ID"],
                row["Municipality English"],
                int(row["Evacuees"]),
                None if pd.isna(temp_value) else int(temp_value),
                f"=ROUNDUP(C{row_number}/50,0)",
                f"=ROUNDUP(C{row_number}/20,0)",
            ]
        )
    for cell in inputs[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    inputs.row_dimensions[1].height = 38
    for row in inputs.iter_rows(min_row=2, max_row=inputs.max_row):
        for cell in row:
            cell.font = Font(name="Aptos", size=8.5, color=text)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = Border(bottom=border)
        for index in (2, 3, 4, 5):
            row[index].alignment = Alignment(horizontal="right", vertical="center")
    for column, width in {"A": 17, "B": 19, "C": 13, "D": 24, "E": 18, "F": 20}.items():
        inputs.column_dimensions[column].width = width
    inputs.freeze_panes = "C2"
    inputs.auto_filter.ref = f"A1:F{inputs.max_row}"

    notes = workbook.create_sheet("Notes")
    notes.sheet_view.showGridLines = False
    notes.column_dimensions["A"].width = 31
    notes.column_dimensions["B"].width = 105
    notes.append(["Field", "Definition"])
    note_rows = [
        ("Pooled requirement", "The ceiling of total city evacuees divided by the benchmark denominator."),
        ("Summed sitewise requirement", "The sum of each shelter's separately rounded-up requirement, including zero requirement at zero-occupancy shelters."),
        ("Fragmentation increment", "Summed sitewise requirement minus pooled requirement; the percentage uses pooled requirement as the denominator."),
        ("Reported temporary toilets", "Observed source total with coverage shown as observed shelter records divided by all shelters. Missing records are retained as missing and never recoded to zero."),
        ("Evidence boundary", "Requirements are benchmark calculations. Reported temporary toilets are not functional capacity because fixed stalls, operability, equipment condition, and toilet-car stall equivalents are not observed consistently."),
    ]
    for row in note_rows:
        notes.append(row)
    for cell in notes[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
    for row in notes.iter_rows(min_row=2, max_row=notes.max_row):
        row[0].font = Font(name="Aptos", size=9, bold=True, color=text)
        row[1].font = Font(name="Aptos", size=9, color=text)
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = Border(bottom=border)
        notes.row_dimensions[row[0].row].height = 45
    notes.freeze_panes = "A2"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)


def validate_output(comparison: pd.DataFrame) -> None:
    workbook = load_workbook(OUTPUT, data_only=False)
    if workbook.sheetnames != ["Requirement Comparison", "Shelter Inputs", "Notes"]:
        raise ValueError("Unexpected workbook sheet structure")
    sheet = workbook["Requirement Comparison"]
    if sheet.max_row != 5 or sheet.max_column != 9:
        raise ValueError("Unexpected main-table dimensions")
    formula_count = sum(
        1
        for workbook_sheet in workbook.worksheets
        for row in workbook_sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )
    if formula_count != 106:
        raise ValueError(f"Unexpected formula count: {formula_count}")
    if any(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for row in sheet.iter_rows()
        for cell in row
    ):
        raise ValueError("The article-facing Requirement Comparison must contain static values")
    japanese_cells = []
    for workbook_sheet in workbook.worksheets:
        for row in workbook_sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and JAPANESE_PATTERN.search(cell.value):
                    japanese_cells.append(f"{workbook_sheet.title}!{cell.coordinate}")
    if japanese_cells:
        raise ValueError(f"Japanese text remains in workbook: {japanese_cells[:5]}")
    observed = comparison.set_index(["Municipality", "Toilet Benchmark"])
    checks = {
        ("Yatsushiro City", "Initial response (1 per 50)"): (38, 56, 18),
        ("Yatsushiro City", "Prolonged stay (1 per 20)"): (93, 112, 19),
        ("Kumamoto City", "Initial response (1 per 50)"): (6, 16, 10),
        ("Kumamoto City", "Prolonged stay (1 per 20)"): (14, 23, 9),
    }
    for key, expected in checks.items():
        row = observed.loc[key]
        actual = tuple(
            int(row[column])
            for column in (
                "Pooled Requirement",
                "Summed Sitewise Requirement",
                "Fragmentation Increment",
            )
        )
        if actual != expected:
            raise ValueError(f"Requirement check failed for {key}: {actual}")


def main() -> None:
    base = prepare_source(pd.read_parquet(SOURCE))
    comparison = expected_comparison(base)
    write_workbook(base, comparison)
    validate_output(comparison)
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print("Main table: 4 municipality-benchmark rows x 9 columns")
    for _, row in comparison.iterrows():
        print(
            f"{row['Municipality']}, {row['Toilet Benchmark']}: "
            f"pooled {int(row['Pooled Requirement'])}, "
            f"sitewise {int(row['Summed Sitewise Requirement'])}, "
            f"increment {int(row['Fragmentation Increment'])}"
        )


if __name__ == "__main__":
    main()
