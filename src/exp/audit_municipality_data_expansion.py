"""Audit municipality shelter totals against official long-term-care tables.

This script writes exploratory evidence audits only. It does not create an
analysis-ready dataset and does not infer shelter-user composition.
"""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SNAPSHOT = ROOT / "data/raw/official_event/kumamoto_prefecture/2026-08-12_1400_municipality_shelter_totals.csv"
LTC_DIR = ROOT / "data/raw/statistics/mhlw_ltc_2026_04"
OUTPUT_DIR = ROOT / "data/exp/data-source-expansion"

LEVEL_COLUMNS = [
    "Support Level 1",
    "Support Level 2",
    "Care Level 1",
    "Care Level 2",
    "Care Level 3",
    "Care Level 4",
    "Care Level 5",
    "Certified Total",
]

LTC_FILES = {
    "Total": LTC_DIR / "2604-h2-1_total.xlsx",
    "Male": LTC_DIR / "2604-h2-2_male.xlsx",
    "Female": LTC_DIR / "2604-h2-3_female.xlsx",
}

EVIDENCE = {
    "熊本市": ("available_near_time_2026-08-12_1200", "available_near_time_2026-08-12_1200", "not_shelter_specific", "partial", "tier_2_shelter_detailed", "municipality screening and shelter-level demand scenarios"),
    "八代市": ("available_near_time_2026-08-12_1200", "available_near_time_2026-08-12_1200", "available_near_time", "available", "tier_2_shelter_detailed", "municipality screening and shelter-level demand and deployment scenarios"),
    "宇土市": ("single_open_shelter_identity_supported", "not_found_at_snapshot", "not_found", "not_found", "tier_1_municipality_plus_identity", "municipality screening; shelter identity as context only"),
    "甲佐町": ("four_open_shelter_identities_supported", "not_found_at_snapshot", "not_found", "not_found", "tier_1_municipality_plus_identity", "municipality screening; shelter identities as context only"),
    "御船町": ("single_shelter_identity_partially_supported", "not_found_at_snapshot", "not_found", "not_found", "tier_1_municipality_plus_partial_identity", "municipality screening only"),
}

DEFAULT_EVIDENCE = (
    "not_found_at_snapshot",
    "not_found_at_snapshot",
    "not_found",
    "not_found",
    "tier_1_municipality_only",
    "municipality screening only",
)


def read_ltc_table(path: Path, sex: str) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = []
    for row in sheet.iter_rows(min_row=7, values_only=True):
        if row[0] != "熊本県" or not row[1]:
            continue
        values = row[2:10]
        record = {"municipality_japanese": str(row[1]).strip()}
        record.update({f"{sex} {name}": value for name, value in zip(LEVEL_COLUMNS, values)})
        rows.append(record)
    return pd.DataFrame(rows)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    snapshot = pd.read_csv(SNAPSHOT, dtype={"municipality_code": "string"})

    merged = snapshot.copy()
    for sex, path in LTC_FILES.items():
        merged = merged.merge(
            read_ltc_table(path, sex),
            on="municipality_japanese",
            how="left",
            validate="one_to_one",
        )

    ltc_columns = [column for column in merged.columns if any(column.startswith(f"{sex} ") for sex in LTC_FILES)]
    if merged[ltc_columns].isna().any().any():
        missing = merged.loc[merged[ltc_columns].isna().any(axis=1), "municipality_japanese"].tolist()
        raise ValueError(f"Unmatched long-term-care records: {missing}")

    total_level_columns = [f"Total {name}" for name in LEVEL_COLUMNS]
    male_level_columns = [f"Male {name}" for name in LEVEL_COLUMNS]
    female_level_columns = [f"Female {name}" for name in LEVEL_COLUMNS]
    if not (merged[total_level_columns].to_numpy() == (merged[male_level_columns].to_numpy() + merged[female_level_columns].to_numpy())).all():
        raise ValueError("Male and female counts do not reproduce published totals")

    audit_columns = [
        "municipality_code",
        "municipality_japanese",
        "municipality_english",
        "observation_time",
        "open_shelters",
        "evacuees",
        *ltc_columns,
    ]
    merged[audit_columns].to_csv(
        OUTPUT_DIR / "mhlw_ltc_11_municipalities_audit.csv",
        index=False,
    )

    coverage = merged[[
        "municipality_code",
        "municipality_japanese",
        "municipality_english",
        "observation_time",
        "open_shelters",
        "evacuees",
    ]].copy()
    coverage["municipality_totals_status"] = "available_exact_time"
    evidence_rows = [EVIDENCE.get(name, DEFAULT_EVIDENCE) for name in coverage["municipality_japanese"]]
    evidence_columns = [
        "shelter_roster_status",
        "shelter_occupancy_status",
        "water_status_availability",
        "deployment_records_availability",
        "evidence_tier",
        "permitted_analysis_level",
    ]
    for index, column in enumerate(evidence_columns):
        coverage[column] = [row[index] for row in evidence_rows]
    coverage["sex_age_input_status"] = "available_unprocessed_2020_mesh"
    coverage["functional_support_input_status"] = "available_observed_2026_04_ltc_by_sex_and_care_level"
    coverage.to_csv(
        OUTPUT_DIR / "municipality_evidence_coverage.csv",
        index=False,
    )

    print(f"Matched municipalities: {len(merged)}")
    print(f"Open shelters: {int(merged['open_shelters'].sum())}")
    print(f"Evacuees: {int(merged['evacuees'].sum())}")
    print(f"Certified total: {int(merged['Total Certified Total'].sum())}")
    print(f"Certified female: {int(merged['Female Certified Total'].sum())}")
    print("Sex reconciliation: exact")
    print(f"Wrote: {OUTPUT_DIR / 'mhlw_ltc_11_municipalities_audit.csv'}")
    print(f"Wrote: {OUTPUT_DIR / 'municipality_evidence_coverage.csv'}")


if __name__ == "__main__":
    main()
